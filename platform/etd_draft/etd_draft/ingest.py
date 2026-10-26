"""元ガイドライン（PDF / DOCX / TXT / MD）を段落ブロックに分解する。

出力は (text, page, source_file) のリスト。CQ への割り当ては segment.py が行う。
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path


@dataclass
class RawBlock:
    text: str
    page: int | None
    source_file: str


_JP_SENTENCE_END = "。．.！!？?」』）)"


def normalize(text: str) -> str:
    """全角英数→半角、空白正規化。CQ 番号検出の安定化のため NFKC を使う。"""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("　", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _lines_to_paragraphs(lines: list[str]) -> list[str]:
    """PDF の行分割を段落に戻す。

    - 空行で段落を区切る
    - 日本語は単語間空白が無いので、行末が文末記号でなければ次行と結合する
    - 行頭が CQ/BQ/FQ 見出しや箇条書き記号なら新段落とする
    """
    paras: list[str] = []
    buf = ""
    heading_re = re.compile(r"^\s*(CQ|BQ|FQ|推奨|解説|エビデンス|背景|文献)\s*[\d０-９]*", re.I)
    bullet_re = re.compile(r"^\s*([・●○◯■□▪➢✓\-*]|\d+[.)．]|[①-⑳])")
    # 「P: …」「推奨の強さ: …」のようなキー:値行は常に独立段落（PICO 抽出のため結合しない）
    kv_re = re.compile(r"^\s*[A-Za-zＡ-Ｚａ-ｚ一-龥ぁ-んァ-ヶー]{1,12}\s*[（(]?[^)）:：]{0,12}[)）]?\s*[:：]\s*\S")
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            if buf:
                paras.append(buf)
                buf = ""
            continue
        if kv_re.match(line):
            if buf:
                paras.append(buf)
                buf = ""
            paras.append(line.strip())
            continue
        if buf and (heading_re.match(line) or bullet_re.match(line)):
            paras.append(buf)
            buf = ""
        if not buf:
            buf = line.strip()
        else:
            # 行末が文末記号なら改行を保持、そうでなければ結合
            if buf[-1] in _JP_SENTENCE_END:
                paras.append(buf)
                buf = line.strip()
            else:
                buf = buf + line.strip()
    if buf:
        paras.append(buf)
    return [p for p in paras if p.strip()]


def _is_noise(text: str) -> bool:
    """ページ番号・柱などのノイズ行。"""
    t = text.strip()
    if len(t) <= 2:
        return True
    if re.fullmatch(r"[-–—\s]*\d{1,4}[-–—\s]*", t):  # ページ番号
        return True
    return False


def read_pdf(path: Path) -> list[RawBlock]:
    import pdfplumber  # 遅延 import

    blocks: list[RawBlock] = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            lines = text.splitlines()
            for para in _lines_to_paragraphs(lines):
                para = normalize(para)
                if not _is_noise(para):
                    blocks.append(RawBlock(para, i, path.name))
    return blocks


def read_docx(path: Path) -> list[RawBlock]:
    import docx  # python-docx

    d = docx.Document(str(path))
    blocks: list[RawBlock] = []
    for p in d.paragraphs:
        t = normalize(p.text)
        if t and not _is_noise(t):
            blocks.append(RawBlock(t, None, path.name))
    # 表のセルも段落として拾う（推奨一覧表が表組みの場合に備える）
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [normalize(c.text) for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                blocks.append(RawBlock(" | ".join(dict.fromkeys(cells)), None, path.name))
    return blocks


def read_text(path: Path) -> list[RawBlock]:
    lines = path.read_text(encoding="utf-8").splitlines()
    return [
        RawBlock(normalize(p), None, path.name)
        for p in _lines_to_paragraphs(lines)
        if not _is_noise(p)
    ]


def read_any(path: str | Path) -> list[RawBlock]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext in (".docx",):
        return read_docx(path)
    if ext in (".txt", ".md", ".markdown"):
        return read_text(path)
    raise ValueError(f"未対応の拡張子: {ext}（pdf / docx / txt / md に対応）")
