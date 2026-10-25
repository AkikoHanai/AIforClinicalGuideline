"""2028年版の推奨作成の記録(EtD)の様式案を、CQごとのシートにする(2023年版の記載を先に入れる)。
  python3 platform/core/build_etd_form2028.py <etd_2023.json> -o <出力.xlsx>

様式案の構成(2023年版の体裁に、2028年版の記入欄を並べる):
  A 推奨(2023年版 | 2028年版案)  B アウトカム別のエビデンス(様式SR-8・SoFから転記)
  C 推奨決定で考慮する項目(2023年版の記載 | 2028年版の判断と根拠)  D 2023年版からの変更点と理由  E 投票(決定記録シートで記録)
判断の選択肢は案であり、統括委員会・SR委員長の承認を受けて確定する。判断は委員が行う。
"""
import argparse, json, os
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PLAN = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "governance", "approved_plan.json"), encoding="utf-8"))
OWNER = {d: "、".join(it["担当"]) for it in PLAN["SR項目"] + PLAN["FRQ項目"] for d in it["dir"]}
STR = {"1": "1 投与(実施)することを強く推奨する", "2": "2 投与(実施)することを提案する", "3": "3 「推奨なし」とする", "4": "4 投与(実施)しないことを提案する", "5": "5 投与(実施)しないことを強く推奨する"}
BAL = "望ましい効果が大きく上回る,望ましい効果がやや上回る,同等,望ましくない効果がやや上回る,望ましくない効果が大きく上回る,判断できない"
VAL = "ばらつきはない,ばらつきがある,判断できない"
RES = "検討が必要(高額が予想される),検討は不要"
THIN = Side(style="thin", color="000000"); BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREY = PatternFill("solid", fgColor="D9D9D9"); YEL = PatternFill("solid", fgColor="F2F2F2")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("etd"); ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()
    items = json.load(open(a.etd, encoding="utf-8")); per = {}
    for it in items: per.setdefault(it["dir"], []).append(it)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws0 = wb.create_sheet("様式案の説明")
    for r in [["2028年版 推奨作成の記録(EtD)様式案"], [""], ["目的", "SR作業(様式4-5、SR-8、SoF)の結果を、推奨作成会議の判断につなぐ。2023年版の体裁(推奨文・強さ・確実性・投票結果・解説)を残し、2028年版の記入欄を並べる"],
              ["記入する人", "SR担当の2名(AとBとCの草案)。会議で投票者が確認・修正する。判断は委員が行う"], ["記入の時期", "2027年4月末(SR-8の完了)から5月上旬(会議の3週間前)"],
              ["入力元", "B: 様式SR-8とSoF(アウトカムごとの確実性、化学療法の層別)。C-1: SR-8の総括。C-2〜C-4: SR担当の案と、会議の議論"],
              ["判断の選択肢", "C-2、C-3、C-4のプルダウンは案。統括委員会・SR委員長の承認を受けて確定する"], ["投票", "決定記録シート(HTML)で記録する。投票者は統括委員と作成グループ委員。有資格者数はCQごとのCOI除外で変わる"],
              ["対象", f"{len(per)}シート(冷却と圧迫は統合。企画書の項目とFRQを含む)"]]: ws0.append(r)
    ws0.column_dimensions["A"].width = 18; ws0.column_dimensions["B"].width = 100; ws0["A1"].font = Font(bold=True, size=13)
    for row in ws0.iter_rows(min_row=3):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BD
        row[0].font = Font(bold=True)
    for d, its in per.items():
        ws = wb.create_sheet(d.replace("・", "")[:31])
        def put(r, vals, bold=False, fill=None):
            for i, v in enumerate(vals):
                c = ws.cell(row=r, column=i + 1, value=v); c.border = BD; c.alignment = Alignment(wrap_text=True, vertical="top")
                if bold: c.font = Font(bold=True)
                if fill is not None: c.fill = fill
        ws.cell(row=1, column=1, value=f"推奨作成の記録(EtD)様式案　{d}").font = Font(bold=True, size=13)
        ws.cell(row=2, column=1, value=f"SR担当: {OWNER.get(d, '')}　／　作成者(記入): ")
        r = 4; put(r, ["A 推奨", "2023年版", "2028年版案(委員が記入)"], True, GREY); r += 1
        for it in its:
            v = it["votes"]
            put(r, ["CQ", it["title"], ""]); r += 1
            put(r, ["推奨文", it["rec"], ""]); r += 1
            put(r, ["推奨の強さ", STR.get(it["strength"], it["strength"]), ""]); st = r; r += 1
            put(r, ["エビデンスの確実性", it["certainty"], ""]); ct = r; r += 1
            if v:
                put(r, ["投票(2023年版)", f"{v['n']}名(棄権: 利益相反{v['coi']}名・SR{v['sr']}名)。1強く推奨 {v['counts'][0]}、2提案する {v['counts'][1]}、3推奨なし {v['counts'][2]}、4しないことを提案 {v['counts'][3]}、5しないことを強く推奨 {v['counts'][4]}(分母{v['denominator']})", "決定記録シートで記録"]); r += 1
            dv = DataValidation(type="list", formula1='"' + ",".join(STR.values()) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(ws.cell(row=st, column=3))
            dv2 = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=True); ws.add_data_validation(dv2); dv2.add(ws.cell(row=ct, column=3))
            r += 1
        put(r, ["B アウトカム別のエビデンス", "(様式SR-8・SoFから転記)", "", "", "", ""], True, GREY); r += 1
        put(r, ["アウトカム", "重要度(1-9)", "研究数(RCT/観察)", "化学療法の層(白金製剤・タキサン系など)ごとの結果", "効果の要約", "確実性(A-D)"], True, YEL); r += 1
        for _ in range(4): put(r, ["", "", "", "", "", ""]); r += 1
        r += 1; put(r, ["C 推奨決定で考慮する項目", "2023年版の記載", "2028年版の判断(案のプルダウン)", "根拠・コメント"], True, GREY); r += 1
        rows = [("C-1 アウトカム全体にわたる総括的なエビデンスの確実性", f"2023年版: 確実性 {its[0]['certainty']}(解説を参照)", "A,B,C,D"),
                ("C-2 望ましい効果と望ましくない効果のバランス", "解説を参照", BAL), ("C-3 患者・市民の価値観と希望", "解説を参照", VAL), ("C-4 資源の利用(コスト)。特に高額が予想される場合のみ検討", "解説を参照", RES)]
        for lab, old, opts in rows:
            put(r, [lab, old, "", ""]); dvx = DataValidation(type="list", formula1='"' + opts + '"', allow_blank=True); ws.add_data_validation(dvx); dvx.add(ws.cell(row=r, column=3)); r += 1
        r += 1; put(r, ["D 2023年版からの変更点と理由", "(新規文献のPMIDを書く)", "", ""], True, GREY); r += 1; put(r, ["", "", "", ""]); r += 2
        put(r, ["E 投票", "決定記録シート(HTML)で記録。ここには、確定した強さと回を転記する", "", ""], True, GREY); r += 1
        put(r, ["確定した推奨の強さ・投票の回・有資格者数", "", "", ""]); r += 2
        put(r, ["参考: 2023年版の解説(原文)", it["narrative"], "", ""], True); ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        for col, w in zip("ABCDEF", [38, 60, 40, 40, 30, 14]): ws.column_dimensions[col].width = w
    wb.save(a.out); print(f"{len(per)}シート → {a.out}")


if __name__ == "__main__":
    main()
