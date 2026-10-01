"""2023年版作成時の委員ワークシート(検索ログ・RoB表・エビデンス表)を minds_review.xlsx に取り込む

  python3 platform/core/import_prior_worksheet.py review_workspace platform/prior_worksheets/*.json

書き込み先:
  検索式            … 「2023年版の検索記録」行(DB・期間・検索式・件数)
  スクリーニングログ … 2023年版の除外/採用内訳を1行のメモとして
  RoB2_Claude下書き  … 各PMID行: デザイン/対照/備考に「2023年版RoB(RoB1)」と結果要約、
                        2×2 が分かる研究は RR と95%CI を「効果:」として追記(フォレストプロットに出る)
既に委員が書いた欄は上書きしない。
"""
import argparse
import glob
import json
import math
import os

from openpyxl import load_workbook

ROB2_COL = {"pmid": 1, "cite": 2, "outcome_id": 3, "design": 4, "comparator": 5, "note": 15}
ROB1_TO_ROB2 = {"ランダム化": "D1", "隠蔽": "D1", "盲検化": "D2", "測定者盲検": "D4", "ITT": "D3", "その他": "D5/その他"}


def rr_ci(a, n1, c, n2):
    if min(a, c) == 0 or a > n1 or c > n2:
        return None
    rr = (a / n1) / (c / n2)
    se = math.sqrt(1 / a - 1 / n1 + 1 / c - 1 / n2)
    return rr, math.exp(math.log(rr) - 1.96 * se), math.exp(math.log(rr) + 1.96 * se)


def _find_row(ws, pmid=None, label=None, year=None):
    """PMID で探し、無ければラベルの姓(どの語でも可)+年(±2年: 電子版と紙版のずれ)で探す"""
    for r in range(2, ws.max_row + 1):
        if pmid and str(ws.cell(row=r, column=1).value or "").strip() == str(pmid):
            return r
    if not label:
        return None
    import re as _re
    words = [w.lower() for w in _re.findall(r"[A-Za-z\u00C0-\u024F][A-Za-z\u00C0-\u024F'\-]{2,}", label) if w.lower() not in ("et", "al")]
    if year is None:
        m = _re.search(r"(19|20)\d\d", label)
        year = int(m.group(0)) if m else None
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(row=r, column=ROB2_COL["cite"]).value or "").lower()
        if not v or not any(w in v for w in words):
            continue
        ym = _re.search(r"(19|20)\d\d", v)
        if year is None or (ym and abs(int(ym.group(0)) - int(year)) <= 2):
            return r
    return None


def _rob2_sheets(wb):
    fixed = ("RoB2_Claude下書き", "RoB2_照合")
    return [wb[n] for n in wb.sheetnames if n.startswith("RoB2_") and n not in fixed]


def _clear_placeholder(ws, r):
    """作業シート末尾の案内行「(新規論文は papers/ に…)」を、行を使う時に消す"""
    c = ws.cell(row=r, column=15)
    if c.value and str(c.value).startswith("(新規論文は"):
        c.value = None


def _first_free(ws):
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value:
            last = r
    return last + 1


def _add_study_row(wb, draft, st):
    """2023年版のSR表にあるが文献リスト(PMID付き)に無い研究を、全RoB2シートに同じ位置で追加"""
    import re as _re
    key = "SR2023:" + _re.sub(r"[^A-Za-z0-9]+", "", st.get("surname") or st.get("label", ""))[:20] + str(st.get("year") or "")
    for ws in [draft] + _rob2_sheets(wb):
        if _find_row(ws, key) is not None:
            continue
        r = _first_free(ws)
        _clear_placeholder(ws, r)
        ws.cell(row=r, column=1).value = key
        ws.cell(row=r, column=ROB2_COL["cite"]).value = st.get("label")
        ws.cell(row=r, column=13).value = "○"      # 2023年版で引用(SR表)
    return _find_row(draft, key)


def _append_note(ws, r, text):
    c = ws.cell(row=r, column=ROB2_COL["note"])
    if c.value and text in str(c.value):
        return
    c.value = (str(c.value) + " / " if c.value else "") + text


def import_one(ws_dir, spec):
    cq_dir = os.path.join(ws_dir, spec["cq_dir"])
    xlsx = os.path.join(cq_dir, "minds_review.xlsx")
    if not os.path.exists(xlsx):
        print(f"  [skip] {xlsx} がありません"); return
    wb = load_workbook(xlsx)
    s = spec.get("search_2023")
    if s and "検索式" in wb.sheetnames:
        ws = wb["検索式"]
        if not any("2023年版の検索記録" in str(ws.cell(row=r, column=1).value or "") for r in range(1, ws.max_row + 1)):
            ws.append([]); ws.append(["2023年版の検索記録", f"{s.get('database')} / 期間 {s.get('period')} / {s.get('hits')}件"])
            ws.append(["2023年版の検索式", s.get("query")])
            ws.append(["→ 今回の検索期間", f"{s.get('period', '').split('–')[-1]} 以降〜現在(前回検索の終了時点から)"])
    ss = spec.get("searches_2023")
    if ss and "検索式" in wb.sheetnames:
        ws = wb["検索式"]
        if not any("2023年版の検索記録" in str(ws.cell(row=r, column=1).value or "") for r in range(1, ws.max_row + 1)):
            ws.append([])
            for i, x in enumerate(ss, 1):
                tag = f"2023年版の検索記録{i}" if len(ss) > 1 else "2023年版の検索記録"
                ws.append([tag, f"{x.get('database') or 'PubMed'} / 期間: {x.get('period_text', '')}"])
                if x.get("query"):
                    ws.append([f"  検索式", x["query"]])
                if x.get("flow_lines"):
                    ws.append(["  件数・除外の流れ", " / ".join(x["flow_lines"])])
            if spec.get("last_search_end_2023"):
                ws.append(["→ 今回(改訂)の検索の起点", f"{spec['last_search_end_2023']} 以降〜現在 ※前回の検索終了日。検索期間の起点は委員会で確認"])
            if spec.get("note"):
                ws.append(["注意", spec["note"]])
    if s and "スクリーニングログ" in wb.sheetnames:
        ws = wb["スクリーニングログ"]
        memo = ("2023年版: " + "、".join(f"{k} {v}件" for k, v in (s.get("excluded") or {}).items()) + " を除外; 採用 "
                + "、".join(f"{k} {v}件" for k, v in (s.get("included") or {}).items()))
        if not any("2023年版:" in str(ws.cell(row=r, column=2).value or "") for r in range(1, ws.max_row + 1)):
            ws.append(["", memo, "2023年版の記録(委員ワークシートより)"])
    draft = wb["RoB2_Claude下書き"] if "RoB2_Claude下書き" in wb.sheetnames else None
    n = 0
    if draft is not None:
        for st in spec.get("studies", []):
            r = _find_row(draft, st.get("pmid"), st.get("label"), st.get("year"))
            if not r:
                if spec.get("add_missing") and (st.get("result_2023") or st.get("rob_2023")):
                    r = _add_study_row(wb, draft, st)
                    _append_note(draft, r, "2023年版のSR表にあり、文献リストにPMIDなし(PMIDを確認して記入)")
                if not r:
                    continue
            for key, col in (("design", "design"), ("comparator", "comparator")):
                if st.get(key) and not draft.cell(row=r, column=ROB2_COL[col]).value:
                    draft.cell(row=r, column=ROB2_COL[col]).value = st[key]
            rob = st.get("rob_2023")
            if rob:
                _append_note(draft, r, "2023年版RoB(RoB1): " + "、".join(f"{k}{v}(→{ROB1_TO_ROB2.get(k, '')})" for k, v in rob.items()))
            if st.get("result_2023"):
                _append_note(draft, r, "2023年版の結果要約: " + st["result_2023"])
            if st.get("n_text") or st.get("drug") or st.get("comparator"):
                _append_note(draft, r, "2023年版の研究表: " + " / ".join(x for x in [
                    f"症例数 {st['n_text']}" if st.get("n_text") else "", f"誘発薬 {st['drug']}" if st.get("drug") else "",
                    f"対照 {st['comparator']}" if st.get("comparator") else ""] if x))
            if st.get("external_quality"):
                _append_note(draft, r, st["external_quality"])
            e = st.get("effect_from_2023")
            if e and st.get("n_int") and st.get("n_ctrl"):
                ci = rr_ci(e["events_int"], st["n_int"], e["events_ctrl"], st["n_ctrl"])
                if ci:
                    rr, lo, hi = ci
                    _append_note(draft, r, f"効果: RR {rr:.2f} ({lo:.2f}–{hi:.2f}) [{e['outcome']}; {e['events_int']}/{st['n_int']} vs {e['events_ctrl']}/{st['n_ctrl']}; {e.get('note', '')}]")
            n += 1
    wb.save(xlsx)
    print(f"{spec['cq_dir']}: {n}研究を取り込み")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workspace_dir")
    ap.add_argument("specs", nargs="+")
    a = ap.parse_args()
    for pat in a.specs:
        for f in glob.glob(pat):
            spec = json.load(open(f, encoding="utf-8"))
            if spec.get("cq_dir") == "*":   # 全CQに対して、ラベルが一致する研究だけ取り込む
                for d in sorted(glob.glob(os.path.join(a.workspace_dir, "CQ*"))):
                    import_one(a.workspace_dir, dict(spec, cq_dir=os.path.basename(d), quiet=True))
            else:
                import_one(a.workspace_dir, spec)


if __name__ == "__main__":
    main()
