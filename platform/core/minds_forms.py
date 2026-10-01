"""Minds 公式様式の定義と読み書き(4-5 評価シート 介入研究 / SR-8 エビデンス総体)

Minds診療ガイドライン作成マニュアル 2020 ver.3.0 の「4-5 評価シート 介入研究」と
「SR-8 評価シート エビデンス総体」の列構成を、そのまま再現する。
- 4-5: アウトカムごとのブロック。各研究を 0(低) / -1(中・疑い) / -2(高) の3段階で評価
- SR-8: アウトカムごとの1行。バイアスリスク/非一貫性/不精確性/非直接性/その他を 0/-1/-2、
        上昇要因を 0/+1/+2、エビデンスの強さ A〜D、重要性 1〜9

公式様式の列(A〜Z)は変更しない。パイプライン用の列は AA 以降に置く:
  AA=キー(PMID等) / AB=コメント / AC=評価指標(使用尺度)
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="EDEDED")
NOTE_FILL = PatternFill("solid", fgColor="F5F5F5")
BOLD = Font(color="000000", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ---- 4-5 の列(A=1 起点) -----------------------------------------------------
# 個別研究のバイアスリスクと非直接性: 全15項目、値は 0 / -1 / -2
ITEMS = [
    ("ランダム化", "選択バイアス"), ("コンシールメント", "選択バイアス"),
    ("盲検化(参加者・医療提供者)", "実行バイアス"), ("盲検化(アウトカム評価者)", "検出バイアス"),
    ("ITT", "症例減少バイアス"), ("アウトカム不完全報告", "症例減少バイアス"),
    ("選択的アウトカム報告", "その他"), ("早期試験中止", "その他"), ("その他のバイアス", "その他"),
    ("バイアスリスク まとめ", "まとめ"),
    ("非直接性 対象", "非直接性"), ("非直接性 介入", "非直接性"), ("非直接性 対照", "非直接性"),
    ("非直接性 アウトカム", "非直接性"), ("非直接性 まとめ", "非直接性"),
]
ITEM_KEYS = ["randomization", "concealment", "blinding_participants", "blinding_assessors", "itt", "incomplete_outcome",
             "selective_reporting", "early_stopping", "other_bias", "bias_summary",
             "ind_population", "ind_intervention", "ind_comparator", "ind_outcome", "ind_summary"]
N_ITEMS = len(ITEMS)
C_CODE, C_DESIGN = 1, 2
C_ITEM0 = 3                       # C列から15項目(C..Q)
C_CTRL_DEN, C_CTRL_NUM, C_CTRL_PCT, C_INT_DEN, C_INT_NUM, C_INT_PCT = 18, 19, 20, 21, 22, 23   # R..W
C_EFF_TYPE, C_EFF_VAL, C_EFF_CI = 24, 25, 26                                                  # X..Z
C_KEY, C_COMMENT, C_INSTR = 27, 28, 29                                                         # AA..AC(パイプライン用)
SPARE_ROWS = 5                    # 新規論文用の予備行(各アウトカムブロックの末尾)
TOP_ROWS = 7                      # ブロック開始前の行数(タイトル〜アウトカム直前)
HEAD_ROWS = 4                     # ブロックの見出し行数
BLOCK_GAP = 2

COL_HEADERS = (["研究コード", "研究デザイン"] + [n for n, _ in ITEMS]
               + ["対照群分母", "対照群分子", "(%)", "介入群分母", "介入群分子", "(%)",
                  "効果指標(種類)", "効果指標(値)", "信頼区間"]
               + ["キー(PMID・事務局用)", "コメント(研究ごと)", "評価指標(使用尺度)"])
COL_WIDTHS = ([24, 10] + [7] * N_ITEMS + [8, 8, 6, 8, 8, 6, 14, 10, 14] + [14, 40, 20])

SCALE_NOTE = ('* 各項目の評価は"高(-2)"、"中/疑い(-1)"、"低(0)"の3段階。'
              "バイアスリスクは 選択・実行・検出・症例減少・その他、非直接性は 対象・介入・対照・アウトカム を評価し、まとめを記入。"
              "化学療法の種類(白金製剤/タキサン系 など)の違いは「非直接性 対象」で評価する(研究特性シート参照)。")

# 化学療法の分類(研究特性シート)。CIPNは薬剤クラスで病態・経過・介入効果が大きく異なる
CHEMO_CLASSES = ["白金製剤", "タキサン系", "ビンカアルカロイド系", "プロテアソーム阻害薬", "複数/混合", "その他", "不明"]
CHEMO_KEYWORDS = [
    ("白金製剤", ["oxaliplatin", "cisplatin", "carboplatin", "platin", "オキサリプラチン", "シスプラチン", "カルボプラチン", "l-ohp", "cddp", "cbdca", "folfox", "xelox", "capox"]),
    ("タキサン系", ["paclitaxel", "docetaxel", "taxane", "taxol", "nab-pac", "abraxane", "パクリタキセル", "ドセタキセル", "タキサン", "ptx", "dtx"]),
    ("ビンカアルカロイド系", ["vincristine", "vinblastine", "vinorelbine", "vinca", "ビンクリスチン", "ビンブラスチン", "ビノレルビン", "vcr", "vbl", "vnr"]),
    ("プロテアソーム阻害薬", ["bortezomib", "carfilzomib", "ボルテゾミブ", "proteasome"]),
]


def classify_chemo(text):
    """薬剤名の文字列から化学療法の分類を推定する(複数クラスなら「複数/混合」)。確定は委員"""
    t = (text or "").lower()
    hit = [c for c, kws in CHEMO_KEYWORDS if any(k in t for k in kws)]
    if not t.strip():
        return "不明"
    if len(hit) == 1:
        return hit[0]
    if len(hit) > 1:
        return "複数/混合"
    return "その他"


def to_score(v):
    """セルの値を 0/-1/-2 に正規化(全角・'−'・文字列も許容)。判定不能は None"""
    if v is None or v == "":
        return None
    s = str(v).strip().replace("−", "-").replace("－", "-").replace("―", "-").replace("ー", "-")
    s = s.translate(str.maketrans("０１２", "012"))
    try:
        n = int(float(s))
    except ValueError:
        return None
    return n if n in (0, -1, -2) else None


def block_len(n_rows):
    return HEAD_ROWS + n_rows + BLOCK_GAP


def block_start(b, n_rows):
    return TOP_ROWS + 1 + b * block_len(n_rows)


def _validation(ws, rng, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "リストから選んでください", "入力値"
    ws.add_data_validation(dv)
    dv.add(rng)


def build_45_sheet(wb, title, cq_title, pico, outcomes, studies, spare=SPARE_ROWS):
    """4-5 評価シート(介入研究)。outcomes: [{id,label}]、studies: [(key, label, design)]"""
    ws = wb.create_sheet(title)
    n_rows = len(studies) + spare
    ws["A1"] = "【4-5　評価シート　介入研究】"
    ws["A1"].font = Font(bold=True, size=12)
    for r, (lab, val) in enumerate([("診療ガイドライン", cq_title), ("対象", pico.get("P") or ""),
                                    ("介入", pico.get("I") or ""), ("対照", pico.get("C") or "")], start=2):
        ws.cell(row=r, column=1, value=lab).font = BOLD
        ws.cell(row=r, column=3, value=val).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    ws.cell(row=3, column=10, value=SCALE_NOTE).alignment = WRAP
    ws.merge_cells(start_row=3, start_column=10, end_row=5, end_column=24)
    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for b, oc in enumerate(outcomes):
        r0 = block_start(b, n_rows)
        ws.cell(row=r0, column=1, value="アウトカム").font = BOLD
        ws.cell(row=r0, column=3, value=oc["label"]).font = BOLD
        ws.cell(row=r0, column=C_KEY, value=oc["id"])
        ws.cell(row=r0, column=C_COMMENT, value=n_rows)      # ブロックの行数(研究行+予備行)。照合の位置計算に使う
        ws.cell(row=r0 + 1, column=1, value="個別研究").font = BOLD
        ws.cell(row=r0 + 1, column=3, value="バイアスリスク*").font = BOLD
        ws.cell(row=r0 + 1, column=11, value="非直接性*").font = BOLD
        ws.cell(row=r0 + 1, column=C_CTRL_DEN, value="リスク人数(アウトカム率)").font = BOLD
        # グループ見出し
        groups = [(3, 4, "選択バイアス"), (5, 5, "実行バイアス"), (6, 6, "検出バイアス"), (7, 8, "症例減少バイアス"),
                  (9, 11, "その他"), (12, 12, "まとめ"), (13, 17, "非直接性*"), (18, 23, "リスク人数(アウトカム率)"),
                  (24, 26, "効果指標")]
        for c0, c1, text in groups:
            ws.cell(row=r0 + 2, column=c0, value=text)
            if c1 > c0:
                ws.merge_cells(start_row=r0 + 2, start_column=c0, end_row=r0 + 2, end_column=c1)
        for c in range(1, C_INSTR + 1):
            ws.cell(row=r0 + 2, column=c).fill = HEADER_FILL
            ws.cell(row=r0 + 2, column=c).font = BOLD
            ws.cell(row=r0 + 2, column=c).alignment = CENTER
            h = ws.cell(row=r0 + 3, column=c, value=COL_HEADERS[c - 1])
            h.fill, h.font, h.alignment, h.border = HEADER_FILL, BOLD, CENTER, BOX
        ws.row_dimensions[r0 + 3].height = 62
        for j in range(n_rows):
            rr = r0 + HEAD_ROWS + j
            if j < len(studies):
                key, label, design = studies[j]
                ws.cell(row=rr, column=C_CODE, value=label)
                ws.cell(row=rr, column=C_DESIGN, value=design or "")
                ws.cell(row=rr, column=C_KEY, value=key)
            for c in range(1, C_INSTR + 1):
                ws.cell(row=rr, column=c).border = BOX
                ws.cell(row=rr, column=c).alignment = WRAP if c in (C_CODE, C_COMMENT, C_INSTR) else CENTER
        first, last = r0 + HEAD_ROWS, r0 + HEAD_ROWS + n_rows - 1
        _validation(ws, f"{get_column_letter(C_ITEM0)}{first}:{get_column_letter(C_ITEM0 + N_ITEMS - 1)}{last}", ["0", "-1", "-2"])
        if b == 0:
            ws.cell(row=first + n_rows, column=1,
                    value="←研究を追加する場合は、上の空き行(研究コードが空の行)を使う。行を挿入する場合は全アウトカムブロックで同じ位置に").font = Font(italic=True)
    ws.freeze_panes = ws.cell(row=TOP_ROWS + 1, column=3)
    return ws


def read_45(ws):
    """4-5シート → ブロックのリスト。各ブロック: {outcome_id, outcome_label, start, rows:[{row,key,label,design,items{},counts{},effect{},instrument,comment}]}"""
    blocks, starts = [], []
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "アウトカム" and ws.cell(row=r, column=C_KEY).value:
            starts.append(r)
    for i, r0 in enumerate(starts):
        end = (starts[i + 1] - 1) if i + 1 < len(starts) else ws.max_row
        rows = []
        for rr in range(r0 + HEAD_ROWS, end + 1):
            code = ws.cell(row=rr, column=C_CODE).value
            key = ws.cell(row=rr, column=C_KEY).value
            if not (code or key):
                continue
            if str(code or "").startswith("←"):
                continue
            g = lambda c: ws.cell(row=rr, column=c).value
            items = {ITEM_KEYS[k]: to_score(g(C_ITEM0 + k)) for k in range(N_ITEMS)}
            counts = {"ctrl_den": g(C_CTRL_DEN), "ctrl_num": g(C_CTRL_NUM), "int_den": g(C_INT_DEN), "int_num": g(C_INT_NUM)}
            eff = {"type": g(C_EFF_TYPE), "value": g(C_EFF_VAL), "ci": g(C_EFF_CI)}
            rows.append({"row": rr, "key": str(key).strip() if key else None, "label": code, "design": g(C_DESIGN),
                         "items": items, "counts": counts, "effect": eff, "instrument": g(C_INSTR), "comment": g(C_COMMENT)})
        nr = ws.cell(row=r0, column=C_COMMENT).value
        blocks.append({"outcome_id": ws.cell(row=r0, column=C_KEY).value, "outcome_label": ws.cell(row=r0, column=3).value,
                       "start": r0, "rows": rows, "n_rows": int(nr) if str(nr).isdigit() else None})
    return blocks


def find_row(ws, outcome_id, key):
    for blk in read_45(ws):
        if blk["outcome_id"] == outcome_id:
            for row in blk["rows"]:
                if row["key"] == str(key):
                    return row["row"]
    return None


def add_study_all_blocks(ws, key, label, design=None):
    """全アウトカムブロックの最初の空き行(研究コードもキーも空)に研究を追加する。
    どのシートでも同じ位置になる(位置で照合するため)。追加できた行数を返す"""
    n = 0
    starts = [r for r in range(1, ws.max_row + 1)
              if ws.cell(row=r, column=1).value == "アウトカム" and ws.cell(row=r, column=C_KEY).value]
    for i, r0 in enumerate(starts):
        end = (starts[i + 1] - 1) if i + 1 < len(starts) else ws.max_row
        if any(str(ws.cell(row=rr, column=C_KEY).value or "") == str(key) for rr in range(r0 + HEAD_ROWS, end + 1)):
            continue
        for rr in range(r0 + HEAD_ROWS, end + 1):
            if not ws.cell(row=rr, column=C_CODE).value and not ws.cell(row=rr, column=C_KEY).value \
                    and ws.cell(row=rr, column=C_CODE).border.left.style == "thin":
                ws.cell(row=rr, column=C_CODE, value=label)
                ws.cell(row=rr, column=C_KEY, value=str(key))
                if design:
                    ws.cell(row=rr, column=C_DESIGN, value=design)
                n += 1
                break
    return n


# ---- SR-8 エビデンス総体 -------------------------------------------------------
SR8_HEADERS = ["アウトカム", "研究デザイン／研究数", "＊バイアスリスク", "＊非一貫性", "＊不精確性", "＊非直接性",
               "＊その他(出版バイアスなど)", "＊＊上昇要因(観察研究)", "対照群分母", "対照群分子", "(％)", "介入群分母", "介入群分子", "(％)",
               "効果指標(種類)", "効果指標統合値", "95%信頼区間", "＊＊＊エビデンスの強さ", "＊＊＊＊重要性", "コメント",
               "絶対効果", "効果指標統合値", "95%信頼区間"]
SR8_NOTE = ("エビデンスの強さについて、RCT は「強(A)」からスタート、観察研究は「弱(C)」からスタート。\n"
            "＊各ドメインは「高(－2)」、「中／疑い(－1)」、「低(0)」の3段階。\n＊＊上昇要因は「高(＋2)」、「中(＋1)」、「低(0)」の3段階。\n"
            "＊＊＊エビデンスの強さは「強(A)」、「中(B)」、「弱(C)」、「非常に弱(D)」の4段階。\n＊＊＊＊重要性はアウトカムの重要性(1〜9)。")
S_KEY, S_STRATUM = 27, 28           # AA=アウトカムID、AB=層別(化学療法の分類)
SR8_COL = {"outcome": 1, "design": 2, "bias": 3, "inconsistency": 4, "imprecision": 5, "indirectness": 6, "other": 7,
           "upgrade": 8, "ctrl_den": 9, "ctrl_num": 10, "int_den": 12, "int_num": 13, "eff_type": 15, "eff_val": 16, "eff_ci": 17,
           "strength": 18, "importance": 19, "comment": 20}


def build_sr8_sheet(wb, title, cq_title, pico, outcomes, strata=None, n_spare=0):
    ws = wb.create_sheet(title)
    ws["A1"] = "【SR-8　評価シート　エビデンス総体(絶対効果指標の結果を記入する場合)】"
    ws["A1"].font = Font(bold=True, size=12)
    for r, (lab, val) in enumerate([("診療ガイドライン", cq_title), ("対象", pico.get("P") or ""),
                                    ("介入", pico.get("I") or ""), ("対照", pico.get("C") or "")], start=2):
        ws.cell(row=r, column=1, value=lab).font = BOLD
        ws.cell(row=r, column=2, value=val).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.cell(row=2, column=9, value=SR8_NOTE).alignment = WRAP
    ws.merge_cells(start_row=2, start_column=9, end_row=5, end_column=20)
    ws.cell(row=7, column=1, value="エビデンス総体").font = BOLD
    ws.cell(row=7, column=9, value="リスク人数(アウトカム率)").font = BOLD
    for c, h in enumerate(SR8_HEADERS, start=1):
        cell = ws.cell(row=8, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = HEADER_FILL, BOLD, CENTER, BOX
    ws.cell(row=8, column=S_KEY, value="アウトカムID(事務局用)").font = BOLD
    ws.cell(row=8, column=S_STRATUM, value="層別(化学療法の分類。全体の行は空)").font = BOLD
    ws.row_dimensions[8].height = 48
    widths = [24, 14, 9, 9, 9, 9, 11, 10, 8, 8, 6, 8, 8, 6, 12, 10, 12, 10, 8, 40, 10, 10, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions[get_column_letter(S_KEY)].width = 16
    ws.column_dimensions[get_column_letter(S_STRATUM)].width = 18
    r = 9
    rows = []
    for oc in outcomes:
        rows.append((oc, None))
        for st in (strata or []):
            rows.append((oc, st))
    for oc, st in rows:
        ws.cell(row=r, column=1, value=(oc["label"] if st is None else f"  └ {oc['label']}(層別: {st})"))
        ws.cell(row=r, column=SR8_COL["importance"], value=oc.get("importance") if st is None else None)
        ws.cell(row=r, column=S_KEY, value=oc["id"])
        if st:
            ws.cell(row=r, column=S_STRATUM, value=st)
            ws.cell(row=r, column=SR8_COL["comment"], value="(任意)層別でエビデンス総体を分ける場合のみ記入。不要なら行を削除")
        for c in range(1, 24):
            ws.cell(row=r, column=c).border = BOX
            ws.cell(row=r, column=c).alignment = WRAP if c in (1, 20) else CENTER
        r += 1
    last = r - 1
    _validation(ws, f"C9:G{last}", ["0", "-1", "-2"])
    _validation(ws, f"H9:H{last}", ["0", "1", "2"])
    _validation(ws, f"R9:R{last}", ["A", "B", "C", "D"])
    ws.cell(row=r + 1, column=1, value="コメント(該当するセルに記入)").font = BOLD
    ws.freeze_panes = "B9"
    return ws


def read_sr8(ws):
    """SR-8 → (総体の行リスト, 層別行リスト)。キー(AA)が空の行は無視"""
    bodies, strata = [], []
    for r in range(9, ws.max_row + 1):
        oid = ws.cell(row=r, column=S_KEY).value
        if not oid:
            continue
        g = lambda k: ws.cell(row=r, column=SR8_COL[k]).value
        def sc(k):
            v = g(k)
            try:
                return int(float(str(v).replace("−", "-").replace("－", "-"))) if v not in (None, "") else None
            except ValueError:
                return None
        row = {"row": r, "outcome_id": str(oid).strip(), "outcome_name": ws.cell(row=r, column=1).value,
               "design_n": g("design"), "bias": sc("bias"), "inconsistency": sc("inconsistency"), "imprecision": sc("imprecision"),
               "indirectness": sc("indirectness"), "other": sc("other"), "upgrade": sc("upgrade"),
               "certainty": (str(g("strength") or "").strip().upper()[:1] or None), "importance": g("importance"),
               "comment": g("comment"), "effect_type": g("eff_type"), "effect_value": g("eff_val"), "effect_ci": g("eff_ci"),
               "stratum": ws.cell(row=r, column=S_STRATUM).value}
        (strata if row["stratum"] else bodies).append(row)
    return bodies, strata


STRENGTH_ORDER = ["A", "B", "C", "D"]
