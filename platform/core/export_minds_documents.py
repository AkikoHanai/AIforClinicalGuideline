"""Mindsに提出・保管する書類の一式を集め、CQごとの作成状況(完了/未完了)を点検する。
  python3 platform/core/export_minds_documents.py <作業ディレクトリ> [--meeting 決定記録_*.json] -o <出力フォルダ>

出力:
  <出力>/CQ別/<CQ>/minds_review.xlsx            検索式・スクリーニング・研究特性・4-5・SR-8・SoF・推奨文草案(CQごとの作業ブック)
  <出力>/推奨作成の記録.xlsx                      推奨一覧 / 投票記録 / 投票除外 / 考慮した項目 / 変更履歴 (会議の決定記録JSONがある場合)
  <出力>/書類の作成状況.xlsx                      CQごと・書類ごとの 完了/一部/未 と、その根拠(件数)
完了の判定は件数による機械的な点検であり、内容の正しさの確認ではない。内容の確認は委員が行う。
"""
import argparse, glob, json, os, shutil
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="000000"); BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STR = {"1": "1 強く推奨", "2": "2 弱く推奨", "3": "3 推奨なし", "4": "4 行わないことを弱く推奨", "5": "5 行わないことを強く推奨"}


def sheet(ws, head, rows, widths=None):
    ws.append(head)
    for r in rows: ws.append(r)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9D9D9")
    for row in ws.iter_rows():
        for c in row: c.border = BD; c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths or [], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def status(done, total):
    if total == 0: return "未"
    return "完了" if done >= total else ("一部" if done > 0 else "未")


def kv(ws):
    return {str(r[0]): r[1] for r in ws.iter_rows(values_only=True) if r and r[0]}


def check_cq(d):
    """1つのCQの作業ブックとcq_packageから、Minds書類ごとの作成状況を数える"""
    out = []
    wb = openpyxl.load_workbook(os.path.join(d, "minds_review.xlsx"), data_only=False)
    pkg = json.load(open(os.path.join(d, "cq_package.json"), encoding="utf-8"))
    frq = pkg.get("question_type") == "FRQ"
    s = kv(wb["検索式"])
    need = ["検索実施日(委員記入)", "検索実施者(委員記入)", "ヒット件数(委員記入)", "重複除去後件数(委員記入)"]
    got = sum(1 for k in need if s.get(k))
    out.append(("検索式と検索結果の記録", status(got, len(need)), f"記入 {got}/{len(need)}項目(検索実施日・実施者・ヒット件数・重複除去後件数)"))
    sc = [r for r in wb["スクリーニングログ"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    s2 = sum(1 for r in sc if r[5]); ex = [r for r in sc if r[5] == "除外" or r[3] == "除外"]
    exr = sum(1 for r in ex if r[4] or r[6])
    out.append(("スクリーニング記録(一次・二次、除外理由)", status(s2, len(sc)) if sc else "未", f"{len(sc)}件中 二次判定済み {s2}件。除外 {len(ex)}件のうち理由あり {exr}件"))
    ch = [r for r in wb["研究特性"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    full = sum(1 for r in ch if r[6] and r[7] and r[8] and r[9] and r[10])
    out.append(("研究特性(化学療法・がん腫・介入・対照)", status(full, len(ch)), f"{len(ch)}件中、化学療法の分類・薬剤・がん腫・介入・対照がすべて記入済み {full}件"))
    if frq:
        out.append(("個別研究の評価(様式4-5)", "対象外(FRQ)", "FRQは推奨を出さないため、必要に応じて記入"))
        out.append(("エビデンス総体(様式SR-8)", "対象外(FRQ)", ""))
        out.append(("SoF表", "対象外(FRQ)", ""))
    else:
        res = pkg.get("results", [])
        fin = [r for r in res if r.get("rob") and r["rob"].get("randomization") is not None and not r.get("provisional")]
        out.append(("個別研究の評価(様式4-5、2名の独立評価・照合済み)", status(len(fin), len(res)) if res else "未", f"評価ブロック {len(res)}件中、2名の評価が一致または確定済み {len(fin)}件(下書きのみは未完了として数える)"))
        eb = pkg.get("evidence_bodies", [])
        ebd = [e for e in eb if e.get("certainty")]
        out.append(("エビデンス総体(様式SR-8)", status(len(ebd), len(eb)) if eb else "未", f"アウトカム {len(eb)}件中、確実性(A〜D)を決定済み {len(ebd)}件"))
        sof = [r for r in wb["SoF"].iter_rows(min_row=2, values_only=True) if r and r[0] and not str(r[0]).startswith("(")]
        out.append(("SoF表", status(len(sof), len(eb)) if eb else "未", f"SoFの行 {len(sof)}件"))
    rec = {str(r[0]): r[1] for r in wb["推奨文草案"].iter_rows(min_row=2, values_only=True) if r and r[0]} if "推奨文草案" in wb.sheetnames else {}
    rt = rec.get("推奨文（草案）") or rec.get("FRQ記載草案") or ""
    out.append(("推奨文草案" if not frq else "FRQ記載草案", "完了" if rt and len(str(rt)) > 10 else "未", "記入あり" if rt else "空"))
    return out


PLAN = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "governance", "approved_plan.json"), encoding="utf-8"))


def prisma_counts(d):
    wb = openpyxl.load_workbook(os.path.join(d, "minds_review.xlsx"))
    s = kv(wb["検索式"])
    sc = [r for r in wb["スクリーニングログ"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    c1 = lambda v: sum(1 for r in sc if r[3] == v); c2 = lambda v: sum(1 for r in sc if r[5] == v)
    reasons = {}
    for r in sc:
        if r[5] == "除外": reasons[r[6] or "(理由未記入)"] = reasons.get(r[6] or "(理由未記入)", 0) + 1
    return [s.get("ヒット件数(委員記入)"), s.get("重複除去後件数(委員記入)"), len(sc), c1("採用"), c1("除外"), c1("保留"), c2("採用"), c2("除外"), "; ".join(f"{k}:{v}" for k, v in reasons.items())]


def ledgers(out):
    """事務局が記入する台帳(COI、外部評価・パブリックコメント、作成経過)。名簿は承認済み企画書"""
    wb = openpyxl.Workbook()
    names = []
    for k in ("統括委員会", "診療ガイドライン作成グループ委員", "SRチーム委員(仮案)", "外部評価委員(仮)"):
        for n in PLAN["体制"][k]:
            names.append([n.split("(")[0], k, "", "", "", ""])
    seen, rows = set(), []
    for r in names:
        if r[0] in seen: continue
        seen.add(r[0]); rows.append(r)
    sheet(wb.active, ["氏名", "役割(企画書)", "COI申告書の提出日", "企業等との関係(有/無)", "関係の内容", "投票除外が必要なCQ(事務局が記入)"], rows, [16, 30, 18, 18, 40, 40])
    wb.active.title = "COI台帳"
    sheet(wb.create_sheet("外部評価・パブコメ対応"), ["No", "種別(外部評価/パブリックコメント)", "意見者", "該当箇所", "意見", "作成グループでの討議日", "変更の要否", "対応(修正内容または変更しない理由)", "対応者"], [], [6, 22, 14, 24, 50, 18, 12, 50, 12])
    sheet(wb.create_sheet("作成経過"), ["項目", "記載内容(診療ガイドライン内に記載する)"], [[k, ""] for k in [
        "作成の目的・経緯(前回: 2023年版)", "作成体制(統括委員会、作成グループ、SRチーム、外部評価委員)", "資金と財源(小林がん学術振興会の助成。資金提供者が内容に関与しないこと)",
        "COIの管理方法", "検索の方法(データベース・期間・検索式)", "文献選択の方法(2名独立、除外理由)", "個別研究の評価方法(様式4-5、0/-1/-2)", "エビデンス総体の評価方法(様式SR-8、A〜D)",
        "推奨作成の方法(投票者・成立条件・回数)", "患者・市民の価値観の反映方法(外部評価委員の関与)", "外部評価・パブリックコメントの方法と対応", "公開の予定(2028年6月頃)と改訂の方針", "AI(Claude)の利用範囲"]], [40, 80])
    wb.save(os.path.join(out, "管理台帳(COI・外部評価・作成経過).xlsx"))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workspace"); ap.add_argument("--meeting"); ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    rows = []
    for d in sorted(glob.glob(os.path.join(a.workspace, "CQ*/"))):
        cqid = os.path.basename(d.rstrip("/"))
        if not os.path.exists(os.path.join(d, "minds_review.xlsx")): continue
        dst = os.path.join(a.out, "CQ別", cqid); os.makedirs(dst, exist_ok=True)
        shutil.copy(os.path.join(d, "minds_review.xlsx"), dst)
        for doc, st, why in check_cq(d): rows.append([cqid, doc, st, why])
    meeting_rows = {}
    if a.meeting:
        m = json.load(open(a.meeting, encoding="utf-8")); S = m["state"]; voters = m["voters"]
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "推奨一覧"
        rec, votes, excl, etd, log = [], [], [], [], []
        ids = [x["id"] for x in m.get("items", [])] or list(S.get("cq", {}))
        for cid in ids:
            c = S.get("cq", {}).get(cid, {})
            fin = c.get("final")
            nv = len(voters); ne = len(c.get("excl") or {})
            rec.append([cid, c.get("rec", ""), STR.get(fin["strength"], "") if fin else "未確定", fin["round"] if fin else "", nv - ne, ne, c.get("certainty", ""), fin["at"][:16].replace("T", " ") if fin else "", c.get("memo", "")])
            for r in ("1", "2", "3"):
                for v, val in (c.get("votes", {}).get(r) or {}).items():
                    votes.append([cid, f"第{r}回", v, STR.get(val, val)])
            for v, why in (c.get("excl") or {}).items(): excl.append([cid, v, why])
            e = c.get("etd") or {}
            etd.append([cid, e.get("benefit", ""), e.get("certainty", ""), e.get("values", ""), e.get("cost", "")])
            meeting_rows[cid] = "確定" if fin else "未確定"
        sheet(ws, ["CQ/FRQ", "推奨文(研究課題)", "推奨の強さ", "確定した投票の回", "有資格者数(分母)", "除外者数", "エビデンスの確実性", "確定日時", "意見の概要"], rec, [34, 60, 22, 10, 10, 8, 10, 18, 40])
        sheet(wb.create_sheet("投票記録"), ["CQ/FRQ", "回", "委員", "投票"], votes, [34, 8, 14, 26])
        sheet(wb.create_sheet("投票除外"), ["CQ/FRQ", "委員", "理由"], excl, [34, 14, 30])
        sheet(wb.create_sheet("考慮した項目"), ["CQ/FRQ", "益と害のバランス", "エビデンス全体の確実性", "患者・市民の価値観と希望", "資源"], etd, [34, 40, 40, 40, 30])
        sheet(wb.create_sheet("変更履歴"), ["日時", "CQ/FRQ", "項目", "値", "記録者"], [[x["t"], x["cq"], x["field"], x["value"], x.get("by", "")] for x in S.get("log", [])], [24, 30, 24, 40, 12])
        sheet(wb.create_sheet("会議"), ["項目", "内容"], [["会議日", S.get("meta", {}).get("date", "")], ["議長", S.get("meta", {}).get("chair", "")], ["記録者", S.get("meta", {}).get("scribe", "")], ["投票者(患者代表は含まない)", "、".join(voters)]], [28, 80])
        wb.save(os.path.join(a.out, "推奨作成の記録.xlsx"))
    for r in rows:
        if r[1].startswith("推奨文草案") or r[1].startswith("FRQ記載草案"):
            r.append(meeting_rows.get(r[0], "会議の記録なし"))
        else:
            r.append("")
    ledgers(a.out)
    pr = []
    for d in sorted(glob.glob(os.path.join(a.workspace, "CQ*/"))):
        if os.path.exists(os.path.join(d, "minds_review.xlsx")): pr.append([os.path.basename(d.rstrip("/"))] + prisma_counts(d))
    wb = openpyxl.Workbook()
    sheet(wb.active, ["CQ/FRQ", "書類", "状況", "根拠(件数)", "会議での決定"], rows, [34, 44, 12, 70, 16])
    wb.active.title = "書類の作成状況"
    sheet(wb.create_sheet("PRISMA件数"), ["CQ/FRQ", "検索ヒット件数", "重複除去後", "スクリーニング対象(ログの件数)", "一次:採用", "一次:除外", "一次:保留", "二次:採用", "二次:除外", "二次の除外理由(件数)"], pr, [34, 10, 10, 14, 8, 8, 8, 8, 8, 50])
    wb.save(os.path.join(a.out, "書類の作成状況.xlsx"))
    tot = {}
    for r in rows: tot[r[2]] = tot.get(r[2], 0) + 1
    print("書類の作成状況:", tot, "→", a.out)


if __name__ == "__main__":
    main()
