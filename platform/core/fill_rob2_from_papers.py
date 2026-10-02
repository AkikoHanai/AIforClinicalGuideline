"""
論文本文 → 4-5_Claude下書き(Minds様式)・研究特性・RoB2(参考) を自動記入
====================================================================
採用論文(2023年版採用 + papers/ の新規PDF)の本文(Method/Result)から Claude (AWS Bedrock) が
  デザイン / 化学療法の分類 / 症例数 / 対照の内容 / アウトカムごとの評価指標 /
  Minds 4-5 のバイアスリスク10項目・非直接性5項目(0/-1/-2) / リスク人数・効果量
を構造化し、次のシートに書き込む。
  4-5_Claude下書き … 様式4-5(アウトカムごとのブロック)。これは参考用の下書き
  研究特性          … 化学療法の分類・症例数・対照など(空欄だけ埋める。委員が確定)
  RoB2(参考)       … Cochrane RoB 2.0 の参考評価(任意)
評価者2名のシート(4-5_<名前>)には書き込まない(Mindsの独立二重評価。下書きに引きずられないため)。

本文の入手順(PDFは必須ではない):
  1. papers/<PMID>.pdf があればそれ(pdftotext)
  2. 無ければ Europe PMC のオープンアクセス全文(PMC)を取得 … 多くの論文はこれで足りる
  3. それも無ければ PubMed の抄録のみ(判定は限定的。要PDFと注記)
取得した本文は papers/<PMID>.txt にキャッシュする(2回目以降は通信しない)。

実行環境: AWS認証情報のあるMac(ネットワーク必要)。
  python3 fill_rob2_from_papers.py <review_workspace> [--only CQ1-冷却・圧迫] [--dry-run] [--no-fetch] [--sync-only]
  --dry-run   : Bedrockを呼ばず、本文取得と書き込み経路だけ確認
  --no-fetch  : ネットワーク取得をしない(PDFとキャッシュのみ)
  --sync-only : papers/ の新規PDFをシート・マニフェストに登録するだけ(=dry-run + no-fetch)
"""
import argparse
import glob
import json
import os
import re
import subprocess
import sys
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from minds_forms import (C_CODE, C_COMMENT, C_CTRL_DEN, C_CTRL_NUM, C_DESIGN, C_EFF_CI, C_EFF_TYPE, C_EFF_VAL, C_INSTR,
                         C_INT_DEN, C_INT_NUM, C_ITEM0, CHEMO_CLASSES, ITEM_KEYS, add_study_all_blocks, classify_chemo, find_row, read_45)

DEFAULT_MODEL = os.environ.get("BEDROCK_MODEL_ID", "anthropic.claude-sonnet-4-5")
DRAFT_SHEET = "4-5_Claude下書き"
FIXED_45 = {DRAFT_SHEET, "4-5_照合"}
KIND_JA = {"pdf": "PDF", "oa_fulltext": "全文(オープンアクセス)", "abstract": "抄録のみ(要PDF)",
           "none": "本文なし(要PDF)"}
UA = {"User-Agent": "AIforClinicalGuideline/1.0 (CIPN guideline revision; contact via repo)"}
CHAR_COL = {"key": 1, "label": 2, "design": 3, "country": 4, "n_total": 5, "n_int": 6, "chemo_class": 7, "chemo_drugs": 8,
            "cancer": 9, "intervention_detail": 10, "comparator_detail": 11, "followup": 12, "fulltext": 13, "note": 14}

SYSTEM = (
    "あなたは診療ガイドライン(Minds診療ガイドライン作成マニュアル2020)の系統的レビュー担当の補助です。"
    "渡された論文本文(主にMethodとResult)だけを根拠に、Minds様式『4-5 評価シート 介入研究』の各項目を"
    "0(低リスク)/-1(中・疑い)/-2(高リスク)で評価し、指定のJSONのみを返してください。"
    "本文に書かれていないことは推測せず null にしてください。評価の最終判断は人間の評価者2名が行います。"
)

PROMPT = """CQ: {cq_title}
対象(P): {population}
介入: {intervention}
アウトカム候補: {outcomes}
本文の種類: {kind}  (抄録のみの場合、判定できる項目だけ埋め、他は null)

評価の基準(Minds 4-5):
- 各項目は 0=低リスク(適切に行われている) / -1=中・疑い(不明・不十分) / -2=高リスク(行われていない)。
- ランダム化=割付順序の生成が適切か。コンシールメント=割付の隠蔽。盲検化(参加者・医療提供者)=実行バイアス。
  盲検化(アウトカム評価者)=検出バイアス。ITT=ITT解析。アウトカム不完全報告=脱落・欠測。
  選択的アウトカム報告・早期試験中止・その他のバイアス。バイアスリスク まとめ=上記を総合した値。
- 非直接性 対象/介入/対照/アウトカム = 本CQのPICOとのずれ(0=一致、-1=やや異なる、-2=大きく異なる)。
  対象では特に「化学療法の種類(白金製剤/タキサン系 等)が本CQの対象と一致するか」を評価する。
- 非RCTの場合、ランダム化・コンシールメントは null とし、デザインを書く。

以下の論文本文からJSONを1つだけ返してください(説明文は不要)。
{{
  "pmid": "本文中にあれば",
  "design": "RCT | meta_analysis | cohort | case_control | cross_sectional | case_series | other",
  "citation": {{"first_author": "第一著者の姓", "year": 西暦, "title": "論文タイトル", "journal": "誌名"}},
  "eligible": true/false  (適格基準: 18歳以上・がん患者・英語・メタ解析/RCT/分析疫学研究),
  "study": {{
    "country": "実施国", "n_total": 数値またはnull, "n_int": 介入群の症例数(数値またはnull),
    "chemo_drugs": "対象患者が受けた化学療法の薬剤名(本文の記載のまま)",
    "chemo_class": "{chemo_classes} のいずれか",
    "cancer": "がん種", "intervention_detail": "介入の内容(用量・期間・頻度)",
    "comparator_detail": "対照の内容(プラセボ/通常ケア/無治療/他の薬など)", "followup": "追跡期間・評価時点"
  }},
  "results": [   // アウトカム候補のIDごとに1要素。本論文が測っていないアウトカムは含めない
    {{
      "outcome_id": "アウトカム候補のID",
      "instrument": "このアウトカムの測定に使った評価指標(例: CTCAE v4.0 grade>=2, EORTC QLQ-CIPN20, FACT-Ntx, NRS, TNS)",
      "items": {{ {item_keys} }},      // 各値は 0 / -1 / -2 / null
      "counts": {{"ctrl_den": 対照群の分母, "ctrl_num": 対照群の分子(イベント数), "int_den": 介入群の分母, "int_num": 介入群の分子}},
      "effect": {{"measure": "RR|OR|HR|MD|SMD|null", "value": 数値またはnull, "ci_low": null, "ci_high": null}},
      "rationale": "判定の根拠(本文の該当箇所を短く。バイアスリスクまとめと非直接性まとめの理由を含める)"
    }}
  ],
  "rob2_reference": {{"D1": "低|懸念あり|高", "D2": "...", "D3": "...", "D4": "...", "D5": "...", "overall": "...", "note": "RCTでなければ理由"}}
}}

--- 論文本文 ---
{text}
"""


# ------------------------------------------------------------------
# 本文の入手
# ------------------------------------------------------------------
def _get(url, timeout=30):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read().decode("utf-8", "replace")


def pdf_text(path, max_chars=60000):
    r = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True)
    t = r.stdout
    m = re.search(r"(?is)\b(methods?|patients and methods|materials and methods)\b", t)
    if m and len(t) > max_chars:
        t = t[m.start():]
    return t[:max_chars]


def europepmc_fulltext_xml_to_text(xml_text, max_chars=60000):
    """JATS XML → Methods/Results 優先のプレーンテキスト"""
    root = ET.fromstring(xml_text)
    def txt(el):
        return " ".join(t.strip() for t in el.itertext() if t and t.strip())
    body = root.find(".//body")
    if body is None:
        return txt(root)[:max_chars]
    secs = body.findall("sec")  # 直下のセクションのみ(入れ子の二重取り防止)
    picked = []
    for sec in secs:
        title = (sec.findtext("title") or "").lower()
        if re.search(r"method|material|patients|design|result|outcome|analysis", title):
            picked.append(txt(sec))
    abstract = root.find(".//abstract")
    out = (txt(abstract) + "\n\n" if abstract is not None else "") + ("\n\n".join(picked) if picked else txt(body))
    return out[:max_chars]


def fetch_fulltext(pmid):
    """Europe PMC: OA全文 → 無ければ PubMed抄録。(text, kind) を返す"""
    q = urllib.parse.quote(f"EXT_ID:{pmid} AND SRC:MED")
    try:
        meta = json.loads(_get(f"https://www.ebi.ac.uk/europepmc/webservices/rest/search?query={q}&format=json&resultType=core"))
        hits = meta.get("resultList", {}).get("result", [])
    except Exception as e:  # noqa: BLE001
        hits = []
    if hits:
        h = hits[0]
        pmcid = h.get("pmcid")
        if pmcid and str(h.get("isOpenAccess", "N")).upper() == "Y":
            try:
                xml_text = _get(f"https://www.ebi.ac.uk/europepmc/webservices/rest/{pmcid}/fullTextXML")
                t = europepmc_fulltext_xml_to_text(xml_text)
                if len(t) > 500:
                    return t, "oa_fulltext"
            except Exception:  # noqa: BLE001
                pass
        abst = h.get("abstractText")
        if abst:
            title = h.get("title", "")
            return f"{title}\n\n{abst}", "abstract"
    # 最後の手段: PubMed efetch の抄録
    try:
        xml_text = _get(f"https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi?db=pubmed&id={pmid}&retmode=xml")
        root = ET.fromstring(xml_text)
        title = root.findtext(".//ArticleTitle") or ""
        abst = " ".join(t.text or "" for t in root.findall(".//AbstractText"))
        if abst:
            return f"{title}\n\n{abst}", "abstract"
    except Exception:  # noqa: BLE001
        pass
    return "", "none"


def get_text_for(pmid, papers_dir, no_fetch):
    """PDF → キャッシュ → ネット取得 の順。(text, kind)"""
    pdf = os.path.join(papers_dir, f"{pmid}.pdf")
    if os.path.exists(pdf):
        return pdf_text(pdf), "pdf"
    cache = os.path.join(papers_dir, f"{pmid}.txt")
    if os.path.exists(cache):
        raw = open(cache, encoding="utf-8").read()
        kind = raw.split("\n", 1)[0].replace("#kind:", "").strip() if raw.startswith("#kind:") else "oa_fulltext"
        return raw.split("\n", 1)[1] if "\n" in raw else "", kind
    if no_fetch:
        return "", "none"
    text, kind = fetch_fulltext(pmid)
    if text:
        os.makedirs(papers_dir, exist_ok=True)
        with open(cache, "w", encoding="utf-8") as f:
            f.write(f"#kind:{kind}\n{text}")
    return text, kind


# ------------------------------------------------------------------
# Bedrock
# ------------------------------------------------------------------
def call_bedrock(system, prompt, model_id):
    import boto3  # 実行はMac側
    client = boto3.client("bedrock-runtime", region_name=os.environ.get("AWS_REGION", "ap-northeast-1"))
    resp = client.converse(
        modelId=model_id,
        system=[{"text": system}],
        messages=[{"role": "user", "content": [{"text": prompt}]}],
        inferenceConfig={"maxTokens": 4000, "temperature": 0},
    )
    text = resp["output"]["message"]["content"][0]["text"]
    m = re.search(r"\{.*\}", text, re.S)
    return json.loads(m.group(0) if m else text)


def extract_one(text, kind, item_ctx, outcomes, model_id, dry_run):
    if dry_run or not text:
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        title = (next((ln for ln in lines if len(ln) > 25), None) or (max(lines, key=len) if lines else ""))[:120]
        return {"pmid": None, "design": None, "eligible": None, "citation": {"first_author": None, "year": None, "title": title, "journal": None},
                "study": {}, "results": [], "rob2_reference": None,
                "_note": (f"dry-run: {KIND_JA.get(kind, kind)} {len(text)}文字。Bedrock未呼び出し" if text
                          else "本文が取得できません。papers/<PMID>.pdf を置いてください")}
    item_keys = ", ".join(f'"{k}": ' for k in ITEM_KEYS)
    prompt = PROMPT.format(cq_title=item_ctx["cq_title"], population=item_ctx.get("population") or "がん薬物療法に伴う末梢神経障害のあるがん患者",
                           intervention=item_ctx["intervention"], outcomes=json.dumps(outcomes, ensure_ascii=False),
                           kind=KIND_JA.get(kind, kind), chemo_classes="/".join(CHEMO_CLASSES), item_keys=item_keys, text=text)
    return call_bedrock(SYSTEM, prompt, model_id)


# ------------------------------------------------------------------
# シートへの書き込み
# ------------------------------------------------------------------
def sheets_45(wb):
    draft = wb[DRAFT_SHEET] if DRAFT_SHEET in wb.sheetnames else None
    reviewers = [wb[n] for n in wb.sheetnames if n.startswith("4-5_") and n not in FIXED_45]
    return draft, reviewers


def chars_index(ws):
    idx = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v:
            idx[str(v).strip()] = r
    return idx


def _num(v):
    try:
        return float(v) if v is not None and v != "" else None
    except (TypeError, ValueError):
        return None


def score(v):
    return v if v in (0, -1, -2) else None


def draft_has_data(ws, key):
    return any(x["key"] == key and (any(v is not None for v in x["items"].values()) or x["effect"]["value"])
               for b in read_45(ws) for x in b["rows"])


def write_45_results(ws, key, res, kind, overwrite=True):
    """4-5下書きの、この研究の行(アウトカムごと)に書き込む。評価者シートには書かない"""
    n = 0
    for rr in res.get("results") or []:
        row = find_row(ws, rr.get("outcome_id"), key)
        if not row:
            continue
        for k, name in enumerate(ITEM_KEYS):
            v = score((rr.get("items") or {}).get(name))
            if v is not None:
                ws.cell(row=row, column=C_ITEM0 + k, value=v)
        c = rr.get("counts") or {}
        for col, k in ((C_CTRL_DEN, "ctrl_den"), (C_CTRL_NUM, "ctrl_num"), (C_INT_DEN, "int_den"), (C_INT_NUM, "int_num")):
            if _num(c.get(k)) is not None:
                ws.cell(row=row, column=col, value=_num(c.get(k)))
        e = rr.get("effect") or {}
        if e.get("measure") and _num(e.get("value")) is not None:
            ws.cell(row=row, column=C_EFF_TYPE, value=e["measure"])
            ws.cell(row=row, column=C_EFF_VAL, value=_num(e["value"]))
            if _num(e.get("ci_low")) is not None and _num(e.get("ci_high")) is not None:
                ws.cell(row=row, column=C_EFF_CI, value=f"{_num(e['ci_low']):g} to {_num(e['ci_high']):g}")
        if rr.get("instrument"):
            ws.cell(row=row, column=C_INSTR, value=rr["instrument"])
        if res.get("design"):
            ws.cell(row=row, column=C_DESIGN, value=res["design"])
        ws.cell(row=row, column=C_COMMENT, value=f"[Claude下書き/本文:{KIND_JA.get(kind, kind)}] {rr.get('rationale') or ''}"[:1000])
        n += 1
    return n


def write_chars(ws, r, res, kind):
    """研究特性: 空欄だけ埋める(委員が確定した値を消さない)。化学療法の分類は薬剤名からも推定する"""
    st = res.get("study") or {}
    cit = res.get("citation") or {}
    def put(name, val):
        if val in (None, ""):
            return
        c = ws.cell(row=r, column=CHAR_COL[name])
        if c.value in (None, ""):
            c.value = val
    put("design", res.get("design"))
    for k in ("country", "n_total", "n_int", "chemo_drugs", "cancer", "intervention_detail", "comparator_detail", "followup"):
        put(k, st.get(k))
    cls = st.get("chemo_class") if st.get("chemo_class") in CHEMO_CLASSES else None
    if not cls and st.get("chemo_drugs"):
        cls = classify_chemo(st["chemo_drugs"])
    put("chemo_class", cls)
    ws.cell(row=r, column=CHAR_COL["fulltext"], value=KIND_JA.get(kind, kind))
    return cit


def write_rob2_ref(wb, key, res):
    if "RoB2(参考)" not in wb.sheetnames or not res.get("rob2_reference"):
        return
    ws = wb["RoB2(参考)"]
    ref = res["rob2_reference"]
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == str(key):
            for j, k in enumerate(["D1", "D2", "D3", "D4", "D5", "overall"]):
                if ref.get(k):
                    ws.cell(row=r, column=3 + j, value=ref[k])
            if ref.get("note"):
                ws.cell(row=r, column=9, value=str(ref["note"])[:500])
            return


def _first_free_row(ws):
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value:
            last = r
    return last + 1


def _append_new_row(wb, key, label, design=None):
    """新規論文: 4-5の全シート(下書き+評価者)の全アウトカムブロックの同じ位置、研究特性、RoB2(参考)に行を追加"""
    draft, reviewers = sheets_45(wb)
    for ws in [draft] + reviewers:
        add_study_all_blocks(ws, key, label, design)
    for name in ("研究特性", "RoB2(参考)"):
        if name in wb.sheetnames:
            ws = wb[name]
            if str(key) in chars_index(ws):
                continue
            r = _first_free_row(ws)
            c = ws.cell(row=r, column=ws.max_column if name == "研究特性" else 9)
            if name == "研究特性" and c.value and str(c.value).startswith("(新規論文は"):
                c.value = None
            ws.cell(row=r, column=1, value=str(key))
            ws.cell(row=r, column=2, value=label)
            if name == "研究特性":
                ws.cell(row=r, column=CHAR_COL["note"], value="新規追加(papers/)")
    return key


def _append_screening(wb, pmid, label, cit):
    if "スクリーニングログ" not in wb.sheetnames:
        return
    ws = wb["スクリーニングログ"]
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "") == str(pmid):
            return
    r = 2
    while r <= ws.max_row and (ws.cell(row=r, column=1).value or ws.cell(row=r, column=2).value):
        if str(ws.cell(row=r, column=1).value or "").startswith("──"):
            break
        r += 1
    ws.insert_rows(r)
    ws.cell(row=r, column=1).value = pmid
    ws.cell(row=r, column=2).value = f"{label}: {cit.get('title') or ''}"[:200]
    ws.cell(row=r, column=3).value = "papers/ に追加(検索/ハンドサーチ)"


def _update_reference_status(wb, status_by_pmid):
    """文献リストの「papers/収集状況」を本文の入手状況で更新"""
    if "文献リスト" not in wb.sheetnames:
        return
    ws = wb["文献リスト"]
    for r in range(2, ws.max_row + 1):
        pmid = str(ws.cell(row=r, column=2).value or "").strip()
        if pmid in status_by_pmid:
            ws.cell(row=r, column=4).value = KIND_JA.get(status_by_pmid[pmid], status_by_pmid[pmid])


def _all_new_items(wb, status_by_pmid):
    """研究特性の「新規追加(papers/)」行すべて(過去の実行分も含む)を MANIFEST 用に集める"""
    titles = {}
    if "スクリーニングログ" in wb.sheetnames:
        ws = wb["スクリーニングログ"]
        for r in range(2, ws.max_row + 1):
            k = str(ws.cell(row=r, column=1).value or "").strip()
            if k:
                titles[k] = str(ws.cell(row=r, column=2).value or "")
    items = []
    if "研究特性" in wb.sheetnames:
        ws = wb["研究特性"]
        for r in range(2, ws.max_row + 1):
            key = str(ws.cell(row=r, column=1).value or "").strip()
            if not key or not str(ws.cell(row=r, column=CHAR_COL["note"]).value or "").startswith("新規追加"):
                continue
            label = str(ws.cell(row=r, column=2).value or "")
            title = titles.get(key, "")
            title = title.split(": ", 1)[1] if title.startswith(label + ": ") else title
            items.append({"pmid": key if key.isdigit() else None, "stem": key, "label": label,
                          "cit": {"title": title}, "kind": status_by_pmid.get(key, "pdf")})
    return items


def _update_manifest(path, new_items, status_by_pmid):
    if not os.path.exists(path):
        return
    txt = open(path, encoding="utf-8").read()
    # 2023年版の行: "| n | [pmid](...) | ☐ | ..." の ☐ を入手状況に
    def repl(m):
        pmid = m.group(2)
        st = status_by_pmid.get(pmid)
        mark = {"pdf": "☑PDF", "oa_fulltext": "☑全文", "abstract": "△抄録", "none": "☐"}.get(st, "☐")
        return f"| {m.group(1)} | [{pmid}]({m.group(3)}) | {mark} |"
    txt = re.sub(r"\| (\d+) \| \[(\d+)\]\(([^)]+)\) \| [☐☑△][^|]* \|", repl, txt)
    head, _, _ = txt.partition("## 新規追加文献(改訂で追加するもの)")
    rows = ["| # | PMID | 収集 | 文献 |", "|---|---|---|---|"]
    for i, it in enumerate(new_items, start=1):
        pmid = it["pmid"]
        link = f"[{pmid}](https://pubmed.ncbi.nlm.nih.gov/{pmid}/)" if pmid else it["stem"]
        cit = it["cit"] or {}
        desc = " ".join(x for x in [it["label"], cit.get("title") or "", cit.get("journal") or ""] if x)[:200]
        mark = {"pdf": "☑PDF", "oa_fulltext": "☑全文", "abstract": "△抄録"}.get(it.get("kind", "pdf"), "☑PDF")
        rows.append(f"| {i} | {link} | {mark} | {desc} |")
    if not new_items:
        rows.append("|  |  | ☐ |  |")
    new = (head + "## 新規追加文献(改訂で追加するもの)\n\n"
           "papers/ に置かれたPDFのうち2023年版に無いもの(fill_rob2_from_papers.py が自動更新)\n\n"
           + "\n".join(rows) + "\n\n凡例: ☑全文=オープンアクセス全文を取得 / ☑PDF=papers/にPDFあり / △抄録=抄録のみ(有料誌。PDFを置いてください) / ☐=未取得\n")
    open(path, "w", encoding="utf-8").write(new)


# ------------------------------------------------------------------
def process_cq(cq_dir, model_id, dry_run, no_fetch):
    xlsx = os.path.join(cq_dir, "minds_review.xlsx")
    pkg_path = os.path.join(cq_dir, "cq_package.json")
    papers = os.path.join(cq_dir, "papers")
    if not os.path.exists(xlsx):
        return None
    pkg = json.load(open(pkg_path, encoding="utf-8")) if os.path.exists(pkg_path) else {}
    outcomes = [{"id": o["id"], "label": o["label"]} for o in pkg.get("outcomes", [])]
    ctx = {"cq_title": pkg.get("title", ""), "intervention": pkg.get("_source", {}).get("intervention", ""),
           "population": (pkg.get("pico") or {}).get("P")}
    wb = load_workbook(xlsx)
    draft, reviewers = sheets_45(wb)
    if draft is None or "研究特性" not in wb.sheetnames:
        print(f"  [skip] {cq_dir}: 4-5_Claude下書き/研究特性 シートがありません(prepare_review_workspace.py --force で作り直してください)")
        return None
    chars = wb["研究特性"]
    idx = chars_index(chars)
    status, new_items = {}, []

    # 1) 新規PDF(2023年版に無いPMID/名前) を登録
    for pdf in sorted(glob.glob(os.path.join(papers, "*.pdf"))):
        stem = os.path.splitext(os.path.basename(pdf))[0]
        if stem in idx:
            continue
        text = pdf_text(pdf)
        res = extract_one(text, "pdf", ctx, outcomes, model_id, dry_run)
        pmid = stem if stem.isdigit() else (str(res.get("pmid")) if res.get("pmid") else None)
        cit = res.get("citation") or {}
        label = (f"{cit.get('first_author')} et al. {cit.get('year')}".strip()
                 if cit.get("first_author") else f"(新規) {stem}")
        key = pmid or stem
        _append_new_row(wb, key, label, res.get("design"))
        _append_screening(wb, key, label, cit)
        idx = chars_index(wb["研究特性"])
        write_chars(wb["研究特性"], idx[key], res, "pdf")
        write_45_results(draft, key, res, "pdf")
        write_rob2_ref(wb, key, res)
        status[key] = "pdf"
        new_items.append({"pmid": pmid, "stem": stem, "label": label, "cit": cit})
        print(f"    [新規] {os.path.basename(pdf)} → {label}")

    # 2) 既存行(2023年版採用)を本文から埋める
    n_done = 0
    for key, r in list(idx.items()):
        if key in status or not key.isdigit():
            continue
        text, kind = get_text_for(key, papers, no_fetch)
        if kind == "none":
            c = chars.cell(row=r, column=CHAR_COL["fulltext"])
            if not c.value:
                c.value = KIND_JA["none"]
                status[key] = kind
            continue
        status[key] = kind
        chars.cell(row=r, column=CHAR_COL["fulltext"], value=KIND_JA.get(kind, kind))
        if dry_run and draft_has_data(draft, key):
            continue  # sync-only/dry-run では、前回(Bedrock)の下書きを仮値で上書きしない
        res = extract_one(text, kind, ctx, outcomes, model_id, dry_run)
        write_chars(chars, r, res, kind)
        n = write_45_results(draft, key, res, kind)
        write_rob2_ref(wb, key, res)
        n_done += 1
        print(f"    {key} [{KIND_JA.get(kind, kind)}] → {res.get('design')} / {n}アウトカム")

    _update_reference_status(wb, status)
    wb.save(xlsx)
    _update_manifest(os.path.join(cq_dir, "MANIFEST.md"), _all_new_items(wb, status), status)
    return {"cq": os.path.basename(cq_dir), "rows": n_done, "new": len(new_items),
            "kinds": {k: sum(1 for v in status.values() if v == k) for k in KIND_JA}}


def main():
    ap = argparse.ArgumentParser(description="論文本文(PDF/OA全文/抄録) → 4-5_Claude下書き・研究特性")
    ap.add_argument("workspace_dir")
    ap.add_argument("--only", help="特定のCQディレクトリ名だけ処理")
    ap.add_argument("--model", default=DEFAULT_MODEL)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-fetch", action="store_true")
    ap.add_argument("--sync-only", action="store_true",
                    help="Bedrockもネット取得もせず、papers/ の新規PDFを登録するだけ")
    args = ap.parse_args()
    dry, nofetch = args.dry_run or args.sync_only, args.no_fetch or args.sync_only

    dirs = sorted(d for d in glob.glob(os.path.join(args.workspace_dir, "*"))
                  if os.path.isdir(d) and not os.path.basename(d).startswith("_"))
    if args.only:
        dirs = [d for d in dirs if os.path.basename(d) == args.only]
    summary = []
    for d in dirs:
        print(f"{os.path.basename(d)}:")
        r = process_cq(d, args.model, dry, nofetch)
        if r:
            summary.append(r)
    if summary:
        tot = {k: sum(r["kinds"].get(k, 0) for r in summary) for k in KIND_JA}
        print(f"\n{len(summary)}CQ / 本文の入手: " + ", ".join(f"{KIND_JA[k]} {v}" for k, v in tot.items() if v))
        need = tot.get("abstract", 0) + tot.get("none", 0)
        if need:
            print(f"→ {need}件は全文が取れていません。文献リスト/MANIFESTで「△抄録」「☐」の論文に papers/<PMID>.pdf を置いてください")
        print("次: merge_rob2_evidence.py → review_bundle.py → render_console.py (run_revision.sh が続けて実行します)")


if __name__ == "__main__":
    main()
