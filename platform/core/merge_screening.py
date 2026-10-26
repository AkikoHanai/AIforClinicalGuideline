"""スクリーニング管理シートの担当者別JSONを照合し、作業ブックのスクリーニング関連シートに反映する。
  python3 platform/core/merge_screening.py <作業ディレクトリ> <JSONのフォルダ> [--min-reviewers 2]

JSONのフォルダ: <フォルダ>/<CQ名>/<CQ名>_<氏名>.screening.json(担当者ごとに最新の1ファイルを使う)
反映先(minds_review.xlsx):
  スクリーニングログ      最終の一次・二次の判定、除外理由、備考(Minds様式の記録)。照合で不一致の項目は「未確定(協議)」
  スクリーニング照合      担当者A・Bの判定を並べ、一致・不一致を表示。協議の結論と理由の欄(委員が記入)。結論が入ると最終判定に使われる
  スクリーニング履歴      判定の変更(変更前・変更後・理由)、文献の追加、全文の入手・読了の全履歴
  スクリーニング進捗      担当者別の作業の進み具合(一次・二次・全文・理由未記入)
判定はAIが行わない。委員の判定を記録し、照合するだけである。"""
import argparse, glob, json, os, sys
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

THIN = Side(style="thin", color="000000"); BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = ["PMID", "タイトル", "出典(検索/ハンドサーチ)", "一次スクリーニング(採用/除外/保留)", "一次除外理由", "二次スクリーニング(採用/除外)", "二次除外理由", "備考"]


def latest_by_reviewer(folder):
    best = {}
    for f in glob.glob(os.path.join(folder, "*.screening.json")):
        try:
            d = json.load(open(f, encoding="utf-8"))
        except Exception:
            continue
        r = d.get("reviewer") or os.path.basename(f)
        if r not in best or (d.get("saved_at") or "") > (best[r].get("saved_at") or ""):
            best[r] = d
    return best


def dec(p, st):
    x = (p.get(st) or {})
    return x.get("decision", ""), x.get("reason", ""), x.get("by", "")


def style(ws, header=True):
    for row in ws.iter_rows():
        for c in row:
            c.border = BD; c.alignment = Alignment(wrap_text=True, vertical="top")
    if header:
        for c in ws[1]:
            c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9D9D9")


def stage_result(ds, concl):
    """ds: 担当者ごとの判定(空は未判定)。concl: 協議の結論。戻り値: (最終, 状態)"""
    if concl:
        return concl, "協議で確定"
    vals = [v for v in ds if v]
    if not vals:
        return "", "未判定"
    if len(vals) < len(ds):
        return vals[0], "片方のみ"
    if len(set(vals)) == 1:
        return vals[0], "一致"
    return "未確定(協議)", "不一致"


def merge_cq(wbpath, folder, min_reviewers):
    revs = latest_by_reviewer(folder)
    names = sorted(revs)
    if not names:
        return None
    wb = openpyxl.load_workbook(wbpath)
    # 既存の協議の記入を引き継ぐ
    old = {}
    if "スクリーニング照合" in wb.sheetnames:
        for r in wb["スクリーニング照合"].iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                old[str(r[0])] = (r[-4] or "", r[-3] or "", r[-2] or "", r[-1] or "")
        del wb["スクリーニング照合"]
    keys, meta = [], {}
    for n in names:
        for p in revs[n]["papers"]:
            if p["key"] not in meta:
                keys.append(p["key"]); meta[p["key"]] = p
    A = names[0]; B = names[1] if len(names) > 1 else None
    pm = {n: {p["key"]: p for p in revs[n]["papers"]} for n in names}
    rows, final = [], []
    for k in keys:
        m = meta[k]; pa = pm[A].get(k, {}); pb = pm[B].get(k, {}) if B else {}
        row = [k, m.get("title", ""), m.get("source", "")]
        res = {}
        for st, lab in (("s1", "一次"), ("s2", "二次")):
            da, ra, ba = dec(pa, st); db, rb, bb = dec(pb, st) if B else ("", "", "")
            direct = bool(m.get("direct")) and st == "s1"
            if direct:
                res[st] = ("", "一次を経ない", da, db, ra, rb); continue
            ds = [da] + ([db] if B else [])
            concl = old.get(k, ("", "", "", ""))[0 if st == "s1" else 2]
            fin, state = stage_result(ds, concl)
            if len(names) < min_reviewers and state == "一致":
                state = "1名のみ"
            res[st] = (fin, state, da, db, ra, rb)
        ft = [str((pm[n].get(k) or {}).get("ft") or "未入手") for n in names]
        row += [res["s1"][2], res["s1"][3], res["s1"][1], res["s2"][2], res["s2"][3], res["s2"][1], ft[0], ft[1] if B else "",
                res["s1"][4] or res["s2"][4], res["s1"][5] or res["s2"][5]]
        o = old.get(k, ("", "", "", ""))
        row += [o[0], o[1], o[2], o[3]]
        rows.append(row)
        reason_final = ""
        if res["s1"][0] == "除外":
            reason_final = res["s1"][4] or res["s1"][5]
        s2reason = (res["s2"][4] or res["s2"][5]) if res["s2"][0] == "除外" else ""
        note = "; ".join(x for x in [m.get("note", ""), f"追加: {m['added_by']}({m.get('added_why','')})" if m.get("added_by") else ""] if x)
        final.append([k, m.get("title", ""), m.get("source", ""), res["s1"][0] if res["s1"][0] else ("" if not m.get("direct") else "—"), o[1] or reason_final,
                      res["s2"][0], o[3] or s2reason, note])
    ws = wb.create_sheet("スクリーニング照合")
    ws.append(["PMID", "題名", "出典", f"一次({A})", f"一次({B or '—'})", "一次の状態", f"二次({A})", f"二次({B or '—'})", "二次の状態", f"全文({A})", f"全文({B or '—'})",
               f"除外理由({A})", f"除外理由({B or '—'})", "一次の協議の結論(委員記入)", "一次の協議の理由", "二次の協議の結論(委員記入)", "二次の協議の理由"])
    for r in rows:
        ws.append(r)
    for col, w in zip("ABCDEFGHIJKLMNOPQ", [10, 40, 18, 10, 10, 10, 10, 10, 10, 8, 8, 28, 28, 14, 24, 14, 24]):
        ws.column_dimensions[col].width = w
    style(ws); ws.freeze_panes = "A2"
    # スクリーニングログ(最終)
    if "スクリーニングログ" in wb.sheetnames:
        lg = wb["スクリーニングログ"]
        lg.delete_rows(2, lg.max_row)
    else:
        lg = wb.create_sheet("スクリーニングログ"); lg.append(HEAD)
    for r in final:
        lg.append(r)
    style(lg)
    for col, w in zip("ABCDEFGH", [11, 50, 22, 18, 30, 18, 30, 40]):
        lg.column_dimensions[col].width = w
    # 履歴
    if "スクリーニング履歴" in wb.sheetnames:
        del wb["スクリーニング履歴"]
    h = wb.create_sheet("スクリーニング履歴"); h.append(["日時", "担当", "PMID", "項目", "変更前", "変更後", "理由"])
    allh = sorted((x for n in names for x in revs[n].get("history", [])), key=lambda x: x.get("t", ""))
    for x in allh:
        h.append([x.get("t", "")[:16].replace("T", " "), x.get("by", ""), x.get("key", ""), x.get("field", ""), x.get("from", ""), x.get("to", ""), x.get("reason", "")])
    style(h)
    for col, w in zip("ABCDEFG", [17, 12, 11, 22, 14, 14, 50]):
        h.column_dimensions[col].width = w
    # 進捗
    if "スクリーニング進捗" in wb.sheetnames:
        del wb["スクリーニング進捗"]
    g = wb.create_sheet("スクリーニング進捗")
    g.append(["担当", "文献数", "一次 判定済", "一次 未判定", "二次 判定済", "二次 対象で未判定", "全文 未入手(二次の対象)", "全文 入手済(未読)", "全文 読了", "除外で理由未記入", "判定の変更回数"])
    for n in names:
        P = revs[n]["papers"]
        d1 = lambda p: dec(p, "s1")[0]; d2 = lambda p: dec(p, "s2")[0]
        need2 = lambda p: (d1(p) in ("採用", "保留") or p.get("direct")) and not d2(p)
        targ = lambda p: d2(p) == "採用" or need2(p)
        g.append([n, len(P), sum(1 for p in P if d1(p)), sum(1 for p in P if not d1(p) and not p.get("direct")), sum(1 for p in P if d2(p)), sum(1 for p in P if need2(p)),
                  sum(1 for p in P if targ(p) and (p.get("ft") or "未入手") == "未入手"), sum(1 for p in P if targ(p) and p.get("ft") == "入手済"), sum(1 for p in P if targ(p) and p.get("ft") == "読了"),
                  sum(1 for p in P if (d1(p) == "除外" and not dec(p, "s1")[1].strip()) or (d2(p) == "除外" and not dec(p, "s2")[1].strip())),
                  sum(1 for x in revs[n].get("history", []) if x.get("reason") not in ("初回の判定", "") and x.get("field") in ("一次", "二次"))])
    style(g)
    wb.save(wbpath)
    conflicts = sum(1 for r in rows if r[5] == "不一致" or r[8] == "不一致")
    return {"reviewers": names, "papers": len(keys), "conflicts": conflicts}


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("workspace"); ap.add_argument("folder"); ap.add_argument("--min-reviewers", type=int, default=2)
    a = ap.parse_args()
    for d in sorted(glob.glob(os.path.join(a.workspace, "CQ*/"))):
        cq = os.path.basename(d.rstrip("/")); x = os.path.join(d, "minds_review.xlsx"); f = os.path.join(a.folder, cq)
        if not os.path.exists(x) or not os.path.isdir(f):
            continue
        r = merge_cq(x, f, a.min_reviewers)
        if r:
            print(f"{cq}: 担当 {'、'.join(r['reviewers'])} / 文献 {r['papers']}件 / 不一致 {r['conflicts']}件")


if __name__ == "__main__":
    main()
