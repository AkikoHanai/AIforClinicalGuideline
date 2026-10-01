"""
minds_review.xlsx の「SR-8_エビデンス総体」シート(Minds様式SR-8) → Minds様式SoF(Summary of
Findings)表を同じworkbookに追記する。
====================================================================
メインリポジトリの sr_minds_formatter.py は「本文抽出→SoF」という別粒度の
入力(extracted.csv、論文単位)を前提にしており、本ワークスペースのCQ単位
(アウトカム総体が最初から1シートにまとまっている)とは形が違うため、この
ワークスペース専用に作り直した軽量版。

前提: 委員がMinds 4.4に沿って「エビデンス総体評価」シートを埋め終えている
こと。埋まっていない行(アウトカムIDが空)は無視するので、未入力の段階で
実行しても空のSoFシートができるだけでエラーにはならない(現状はこの状態)。

使い方:
  python3 build_sof.py <review_workspaceディレクトリ>
  (配下の */minds_review.xlsx を全部処理し、SoFシートを追記/更新する)
"""
import argparse
import glob
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from minds_forms import read_sr8

# 印刷時のインク消費を抑えるため、濃い塗り＋白抜き文字は使わない(白地に黒文字)
HEADER_FILL = PatternFill("solid", fgColor="EDEDED")
HEADER_FONT = Font(color="000000", bold=True)
CERT_FILL = {
    "A": PatternFill("solid", fgColor="D4EDDA"),
    "B": PatternFill("solid", fgColor="D1ECF1"),
    "C": PatternFill("solid", fgColor="FFF3CD"),
    "D": PatternFill("solid", fgColor="F8D7DA"),
}
CERT_SYMBOL = {
    "A": "⊕⊕⊕⊕ 強",
    "B": "⊕⊕⊕○ 中",
    "C": "⊕⊕○○ 弱",
    "D": "⊕○○○ 非常に弱い",
}
WRAP = Alignment(wrap_text=True, vertical="top")

SOF_HEADERS = ["アウトカム", "重要度(1-9)", "研究数", "エビデンスの確実性(GRADE)",
               "格下げ理由", "効果の要約", "コメント(委員記入)"]
SOF_WIDTHS = [22, 10, 10, 20, 24, 46, 30]


DOMAIN_JA = {"bias": "バイアスリスク", "inconsistency": "非一貫性", "imprecision": "不精確性",
             "indirectness": "非直接性", "other": "その他(出版バイアス等)"}


def read_evidence_rows(ws):
    """SR-8シートから、アウトカムIDが入っている行(層別行を含む)を拾う"""
    bodies, strata = read_sr8(ws)
    rows = []
    for b in bodies:
        rows.append(b)
        rows.extend(x for x in strata if x["outcome_id"] == b["outcome_id"] and (x["certainty"] or x.get("bias") is not None))
    return rows


def build_sof(wb):
    if "SR-8_エビデンス総体" not in wb.sheetnames:
        return None, 0
    src = wb["SR-8_エビデンス総体"]
    rows = read_evidence_rows(src)

    if "SoF" in wb.sheetnames:
        del wb["SoF"]
    ws = wb.create_sheet("SoF")
    for i, h in enumerate(SOF_HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = WRAP
        ws.column_dimensions[get_column_letter(i)].width = SOF_WIDTHS[i - 1]
    ws.freeze_panes = "A2"

    r, n_filled = 2, 0
    for b in rows:
        cert = (b["certainty"] or "").strip().upper()
        if not cert and not any(b.get(k) is not None for k in DOMAIN_JA):
            continue                       # 未入力の行はSoFに出さない
        n_filled += 1
        reasons = ", ".join(f"{DOMAIN_JA[k]}({b[k]})" for k in DOMAIN_JA if b.get(k) is not None and b[k] < 0) \
            or ("-" if cert == "A" else "(未記入)")
        eff = " ".join(str(x) for x in (b.get("effect_type"), b.get("effect_value")) if x) + \
            (f" (95%CI {b['effect_ci']})" if b.get("effect_ci") else "")
        name = b["outcome_name"] if not b.get("stratum") else f"  └ 層別: {b['stratum']}"
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=b.get("importance"))
        ws.cell(row=r, column=3, value=b.get("design_n"))
        ccell = ws.cell(row=r, column=4, value=CERT_SYMBOL.get(cert, cert))
        if cert in CERT_FILL:
            ccell.fill = CERT_FILL[cert]
        ws.cell(row=r, column=5, value=reasons)
        ws.cell(row=r, column=6, value=eff)
        ws.cell(row=r, column=7, value=b.get("comment") or "")
        for col in (1, 5, 6, 7):
            ws.cell(row=r, column=col).alignment = WRAP
        r += 1

    if not n_filled:
        ws.cell(row=2, column=1,
                value="(SR-8 エビデンス総体シートが未記入のため、SoF表はまだ空です。"
                      "委員がMinds 4.4に沿って記入すると次回実行時に自動反映されます)")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SOF_HEADERS))
        ws.cell(row=2, column=1).alignment = WRAP
    return ws, n_filled


def main():
    ap = argparse.ArgumentParser(description="エビデンス総体評価 → SoF表 生成")
    ap.add_argument("workspace_dir", help="review_workspace/ (CQ*/minds_review.xlsx を含むディレクトリ)")
    args = ap.parse_args()

    files = sorted(glob.glob(os.path.join(args.workspace_dir, "*", "minds_review.xlsx")))
    total_rows = 0
    for fp in files:
        wb = load_workbook(fp)
        ws, n = build_sof(wb)
        if ws is None:
            print(f"[skip] {fp}: SR-8_エビデンス総体 シートが見つかりません")
            continue
        wb.save(fp)
        total_rows += n
        status = f"{n}アウトカム" if n else "未記入(空のSoFシートのみ作成)"
        print(f"{os.path.basename(os.path.dirname(fp)):<28} → SoF: {status}")

    print(f"\n{len(files)}件のworkbookを処理。SoFに反映された総アウトカム数: {total_rows}")


if __name__ == "__main__":
    main()
