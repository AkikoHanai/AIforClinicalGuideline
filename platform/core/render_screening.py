"""スクリーニング管理シート(HTML 1ファイル/CQ)を作る。
  python3 platform/core/render_screening.py <作業ディレクトリ> -o <出力フォルダ>

管理できる作業: 一次スクリーニング(タイトル・抄録)の判定と進捗 / 二次スクリーニング(全文)の対象の追加 /
全文の入手と読了 / 判定の変更(理由が必須。履歴に残る) / 除外理由の記録 / 作業の一覧(未判定・全文未入手・全文未読・理由未記入・保留)。
2名が別々に判定する。担当者ごとに JSON を書き出し、merge_screening.py が照合して作業ブックのスクリーニングログに反映する。"""
import argparse, glob, json, os, re
import openpyxl

HTML = r"""<!doctype html><html lang="ja"><head><meta charset="utf-8"><title>スクリーニング管理 __CQ__</title>
<style>
body{font-family:"Yu Gothic","Hiragino Sans",sans-serif;color:#000;background:#fff;margin:0;font-size:13.5px}
header{padding:8px 16px;border-bottom:2px solid #000}h1{font-size:17px;margin:0 0 4px}
.bar{padding:6px 16px;border-bottom:1px solid #000}.chips button,.btn{border:1px solid #000;background:#fff;color:#000;padding:3px 10px;margin:2px 4px 2px 0;border-radius:11px;cursor:pointer;font-size:12.5px}
.chips button.on{font-weight:bold;border-width:3px}.btn{border-radius:3px}
table{border-collapse:collapse;width:100%}th,td{border:1px solid #000;padding:3px 5px;vertical-align:top;font-size:12.5px}th{background:#eee;position:sticky;top:0}
select,input,textarea{font:inherit;border:1px solid #000;padding:2px;box-sizing:border-box}textarea{width:100%}
.note{font-size:12px}.sm{font-size:11.5px}.hide{display:none}dialog{border:2px solid #000;max-width:760px}
.cnt{display:inline-block;border:1px solid #000;padding:1px 8px;margin-right:6px}
@media print{.bar,.noprint{display:none}}
</style></head><body>
<header><h1>スクリーニング管理　__CQ__　<span class="note">__TITLE__</span></h1>
<div class="noprint">担当者名 <input id="who" size="10"> <button class="btn" onclick="exportJson()">自分の判定をJSONで書き出す</button>
<label class="btn">JSONを読み込む<input type="file" id="imp" accept=".json" class="hide"></label>
<button class="btn" onclick="addDlg.showModal()">文献を追加</button> <button class="btn" onclick="impDlg.showModal()">検索結果を取り込む</button>
<span class="note">入力はこの端末のブラウザに自動保存。共有はJSONの受け渡し。判定の変更は理由が必要で、履歴に残る。</span></div></header>
<div class="bar" id="counts"></div>
<div class="bar chips" id="chips"></div>
<div style="overflow:auto"><table><thead><tr><th>PMID</th><th>文献</th><th>出典</th><th>一次(タイトル・抄録)</th><th>二次(全文)</th><th>全文</th><th>備考・履歴</th></tr></thead><tbody id="rows"></tbody></table></div>
<dialog id="addDlg"><b>文献を追加</b><p class="note">二次スクリーニングの対象として追加する場合(既存ガイドラインの引用文献、ハンドサーチ、作成委員のサーベイランスなど)は、「一次を経ない」を選ぶ。</p>
<p>PMID <input id="a_pmid" size="12"> 題名 <input id="a_title" size="50"></p>
<p>出典 <select id="a_src"><option>検索(PubMed)</option><option>更新検索</option><option>既存ガイドラインの引用文献</option><option>ハンドサーチ</option><option>作成委員のサーベイランス</option><option>その他</option></select>
<label><input type="checkbox" id="a_direct"> 一次を経ない(二次の対象として追加)</label></p>
<p>追加の理由 <input id="a_why" size="60"></p><p><button class="btn" onclick="addPaper()">追加</button> <button class="btn" onclick="addDlg.close()">閉じる</button></p></dialog>
<dialog id="impDlg"><b>検索結果を取り込む</b><p class="note">PubMedの「Save → CSV」の内容、または「PMID(タブ)題名」の行を貼り付ける。すでにある文献は追加しない。</p>
<textarea id="i_text" rows="9"></textarea><p>出典 <select id="i_src"><option>検索(PubMed)</option><option>更新検索</option></select> <button class="btn" onclick="importText()">取り込む</button> <button class="btn" onclick="impDlg.close()">閉じる</button></p></dialog>
<dialog id="hisDlg"><b>履歴</b><div id="hisBody"></div><p><button class="btn" onclick="hisDlg.close()">閉じる</button></p></dialog>
<script>
const CQ=__CQJ__, INIT=__INIT__;
const REASONS=["18歳未満が対象","がん患者でない","英語でない","国内で施行できない方法・薬物","研究デザインが対象外(総説・症例報告など)","CIPNに関する研究でない","本CQの介入を評価していない","予防・治療の目的が本CQと異なる","対象の化学療法が本CQと異なる","重複","計画書のみ","副次解析","主要評価が本CQのアウトカムでない","全文が入手できない","その他(備考に記載)"];
const KEY="cipn-screen-"+CQ; let S={who:"",papers:INIT.map(x=>({...x})),history:[]}; let filt="all";
try{const r=localStorage.getItem(KEY); if(r){S=JSON.parse(r);}}catch(e){}
const $=id=>document.getElementById(id); const esc=s=>String(s==null?"":s).replace(/&/g,"&amp;").replace(/</g,"&lt;").replace(/>/g,"&gt;").replace(/"/g,"&quot;").replace(/'/g,"&#39;");
function save(){try{localStorage.setItem(KEY,JSON.stringify(S));}catch(e){}}
function now(){return new Date().toISOString();}
function who(){return $("who").value.trim()||"(未入力)";}
function d1(p){return (p.s1||{}).decision||"";} function d2(p){return (p.s2||{}).decision||"";}
function pend1(p){return !d1(p) && !p.direct;} function need2(p){return (d1(p)==="採用"||d1(p)==="保留"||p.direct) && !d2(p);}
function noReason(p){return (d1(p)==="除外"&&!(p.s1.reason||"").trim())||(d2(p)==="除外"&&!(p.s2.reason||"").trim());}
const FILTERS={all:["すべて",p=>true],p1:["一次 未判定",pend1],p2:["二次 未判定",need2],ft0:["全文 未入手",p=>(d2(p)==="採用"||need2(p))&&(p.ft||"未入手")==="未入手"],ft1:["全文 未読",p=>(d2(p)==="採用"||need2(p))&&p.ft==="入手済"],nr:["理由 未記入",noReason],hold:["一次 保留",p=>d1(p)==="保留"],added:["追加した文献",p=>!!p.added_by]};
function chips(){$("chips").innerHTML=Object.entries(FILTERS).map(([k,[l,f]])=>`<button class="${k===filt?'on':''}" onclick="filt='${k}';draw()">${l} ${S.papers.filter(f).length}</button>`).join("");}
function counts(){const P=S.papers,c=(f)=>P.filter(f).length;
 $("counts").innerHTML=`<span class="cnt">文献 ${P.length}</span><span class="cnt">一次 判定済 ${c(p=>d1(p))}/${c(p=>!p.direct)}(採用 ${c(p=>d1(p)==="採用")}・除外 ${c(p=>d1(p)==="除外")}・保留 ${c(p=>d1(p)==="保留")})</span><span class="cnt">二次 判定済 ${c(p=>d2(p))}(採用 ${c(p=>d2(p)==="採用")}・除外 ${c(p=>d2(p)==="除外")})/ 対象 ${c(p=>d2(p)||need2(p))}</span><span class="cnt">全文 入手済・読了 ${c(p=>(d2(p)||need2(p))&&(p.ft==="入手済"||p.ft==="読了"))}/${c(p=>d2(p)||need2(p))}</span><span class="cnt">理由未記入 ${c(noReason)}</span>`;}
function reasonSel(p,st){const cur=(p[st]||{}).reason||"";return `<select onchange="setReason('${esc(p.key)}','${st}',this.value)"><option value="">(理由)</option>${REASONS.map(r=>`<option ${r===cur?'selected':''}>${r}</option>`).join("")}</select>`;}
function dec(p,st,opts){const cur=(p[st]||{}).decision||""; const dis=(st==="s2"&&!(d1(p)==="採用"||d1(p)==="保留"||p.direct))?"disabled":"";
 return opts.map(o=>`<label style="margin-right:6px"><input type="radio" name="${st}_${esc(p.key)}" ${cur===o?'checked':''} ${dis} onclick="setDec('${esc(p.key)}','${st}','${o}')">${o}</label>`).join("")+(cur==="除外"?"<br>"+reasonSel(p,st):"")+(cur?`<div class="sm">${esc((p[st]||{}).by||"")} ${esc(((p[st]||{}).at||"").slice(0,10))}</div>`:"");}
function draw(){chips();counts();const f=FILTERS[filt][1];
 $("rows").innerHTML=S.papers.filter(f).map(p=>`<tr><td>${/^\d+$/.test(p.key)?`<a href="https://pubmed.ncbi.nlm.nih.gov/${p.key}/" target="_blank" rel="noopener">${p.key}</a>`:esc(p.key)}</td><td>${esc(p.title||"(題名未入力)")}<div class="sm">${esc(p.year||"")} ${esc(p.author||"")}</div></td><td class="sm">${esc(p.source||"")}${p.added_by?`<br>追加: ${esc(p.added_by)}<br>${esc(p.added_why||"")}`:""}${p.direct?"<br>一次を経ない":""}</td>
 <td>${p.direct?"—":dec(p,"s1",["採用","除外","保留"])}</td><td>${dec(p,"s2",["採用","除外"])}</td>
 <td><select onchange="setFt('${esc(p.key)}',this.value)">${["未入手","入手済","読了"].map(o=>`<option ${((p.ft||"未入手")===o)?'selected':''}>${o}</option>`).join("")}</select></td>
 <td><textarea rows="2" oninput="setNote('${esc(p.key)}',this.value)">${esc(p.note||"")}</textarea><button class="btn" onclick="showHis('${esc(p.key)}')">履歴 ${S.history.filter(h=>h.key===p.key).length}</button></td></tr>`).join("");save();}
function P(k){return S.papers.find(p=>p.key===k);}
function log(k,field,from,to,reason){S.history.push({t:now(),key:k,field,from,to,reason,by:who()});}
function setDec(k,st,val){const p=P(k); const cur=(p[st]||{}).decision||"";
 if(cur===val){return;} let reason="";
 if(cur){reason=window.prompt("判定を変更します("+cur+" → "+val+")。変更の理由を入力してください。"); if(!reason||!reason.trim()){draw();return;}}
 p[st]=Object.assign(p[st]||{}, {decision:val,at:now(),by:who()}); if(val!=="除外"){p[st].reason="";}
 log(k,st==="s1"?"一次":"二次",cur||"未判定",val,reason||"初回の判定"); draw();}
function setReason(k,st,v){const p=P(k); const old=(p[st]||{}).reason||""; p[st].reason=v; log(k,(st==="s1"?"一次":"二次")+"の除外理由",old,v,"理由の記録"); draw();}
function setFt(k,v){const p=P(k); const old=p.ft||"未入手"; if(old===v)return; p.ft=v; log(k,"全文",old,v,""); draw();}
function setNote(k,v){P(k).note=v; save();}
function showHis(k){const h=S.history.filter(x=>x.key===k); $("hisBody").innerHTML=h.length?`<table><tr><th>日時</th><th>担当</th><th>項目</th><th>変更前</th><th>変更後</th><th>理由</th></tr>`+h.map(x=>`<tr><td>${esc(x.t.slice(0,16).replace("T"," "))}</td><td>${esc(x.by)}</td><td>${esc(x.field)}</td><td>${esc(x.from)}</td><td>${esc(x.to)}</td><td>${esc(x.reason)}</td></tr>`).join("")+"</table>":"履歴はありません"; hisDlg.showModal();}
function addPaper(){const k=$("a_pmid").value.trim(); if(!k){alert("PMIDを入力してください");return;} if(P(k)){alert("すでにある文献です");return;} const why=$("a_why").value.trim(); if(!why){alert("追加の理由を入力してください");return;}
 const direct=$("a_direct").checked; S.papers.push({key:k,title:$("a_title").value.trim(),source:$("a_src").value,added_by:who(),added_at:now(),added_why:why,direct,s1:{},s2:{},ft:"未入手"}); log(k,"追加","",direct?"二次の対象として追加":"一次の対象として追加",why); addDlg.close(); draw();}
function importText(){const lines=$("i_text").value.split(/\r?\n/).filter(x=>x.trim()); let n=0;
 for(const l of lines){const m=l.match(/^\s*"?(\d{5,9})"?\s*[\t,]\s*"?(.*?)"?\s*$/)||l.match(/^\s*"?(\d{5,9})"?\s*$/); if(!m)continue; if(P(m[1]))continue; S.papers.push({key:m[1],title:(m[2]||"").slice(0,300),source:$("i_src").value,s1:{},s2:{},ft:"未入手"}); log(m[1],"取り込み","",$("i_src").value,"検索結果"); n++;}
 alert(n+"件を追加しました"); impDlg.close(); draw();}
function exportJson(){const a=document.createElement("a"); a.href=URL.createObjectURL(new Blob([JSON.stringify({cq:CQ,reviewer:who(),saved_at:now(),papers:S.papers,history:S.history},null,1)],{type:"application/json"})); a.download=CQ+"_"+who()+".screening.json"; a.click();}
$("imp").addEventListener("change",e=>{const f=e.target.files[0]; if(!f)return; const r=new FileReader(); r.onload=()=>{try{const d=JSON.parse(r.result); if(!confirm("この端末の入力を、読み込んだ内容で置き換えます。"))return; S={who:d.reviewer,papers:d.papers,history:d.history||[]}; $("who").value=d.reviewer||""; save(); draw();}catch(x){alert("読み込めません: "+x);}}; r.readAsText(f);});
$("who").value=S.who||""; $("who").addEventListener("input",()=>{S.who=$("who").value; save();}); draw();
</script></body></html>"""


def initial(ws):
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        key = str(r[0]).strip(); title = str(r[1] or "")
        m = re.match(r"\s*([^:：]+?)\s*(\d{4})", title)
        p = {"key": key, "title": title, "source": str(r[2] or ""), "year": m.group(2) if m else "", "author": "", "ft": "未入手", "s1": {}, "s2": {}}
        if len(r) > 3 and r[3]:
            p["s1"] = {"decision": str(r[3]), "reason": str(r[4] or ""), "at": "", "by": "既定(作業ブック)"}
        if len(r) > 5 and r[5]:
            p["s2"] = {"decision": str(r[5]), "reason": str(r[6] or ""), "at": "", "by": "既定(作業ブック)"}
        if len(r) > 7 and r[7]:
            p["note"] = str(r[7])
        out.append(p)
    return out


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("workspace"); ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args(); os.makedirs(a.out, exist_ok=True); n = 0
    for d in sorted(glob.glob(os.path.join(a.workspace, "CQ*/"))):
        x = os.path.join(d, "minds_review.xlsx")
        if not os.path.exists(x):
            continue
        cq = os.path.basename(d.rstrip("/")); wb = openpyxl.load_workbook(x, read_only=True)
        init = initial(wb["スクリーニングログ"]) if "スクリーニングログ" in wb.sheetnames else []
        title = ""
        if "CQ・PICO" in wb.sheetnames:
            kv = {str(r[0]): r[1] for r in wb["CQ・PICO"].iter_rows(values_only=True) if r and r[0]}
            title = str(kv.get("CQ(分類)", "")) + " " + str(kv.get("介入/薬剤", ""))
        html = HTML.replace("__CQJ__", json.dumps(cq, ensure_ascii=False)).replace("__INIT__", json.dumps(init, ensure_ascii=False)).replace("__CQ__", cq).replace("__TITLE__", title)
        open(os.path.join(a.out, f"{cq}.screening.html"), "w", encoding="utf-8").write(html); n += 1
    print(f"{n}件のスクリーニング管理シートを作成 → {a.out}")


if __name__ == "__main__":
    main()
