"""Minds 4-5 / SR-8 様式の読み書き・照合・マージのテスト"""
import json
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "core"))

from openpyxl import Workbook, load_workbook  # noqa: E402

import minds_forms as mf  # noqa: E402
import merge_rob2_evidence as mg  # noqa: E402

OUTS = [{"id": "O:A", "label": "CIPN発症頻度", "importance": 9}, {"id": "O:B", "label": "症状の軽減", "importance": 9}]
STUDIES = [("111", "Smith et al. 2013", "RCT"), ("222", "Kono et al. 2013", "RCT")]


def make_wb():
    wb = Workbook()
    wb.remove(wb.active)
    pico = {"P": "がん患者", "I": "X", "C": "プラセボ"}
    for n in ("4-5_Claude下書き", "4-5_甲", "4-5_乙"):
        mf.build_45_sheet(wb, n, "CQ", pico, OUTS, STUDIES)
    n_rows = len(STUDIES) + mf.SPARE_ROWS
    # 照合シート(prepare と同じ並び)
    ws = wb.create_sheet("4-5_照合")
    ws.append(["a", "b", "c", "d", "e", "f", "g"])
    for _ in OUTS:
        for _ in range(n_rows):
            for _ in mf.ITEMS:
                ws.append([None] * 7)
    mf.build_sr8_sheet(wb, "SR-8_エビデンス総体", "CQ", pico, OUTS, strata=["白金製剤"])
    return wb, n_rows


def test_score_normalization():
    assert mf.to_score("－1") == -1 and mf.to_score("ー2") == -2 and mf.to_score(0) == 0
    assert mf.to_score("?") is None and mf.to_score(1) is None and mf.to_score("") is None


def test_chemo_classification():
    assert mf.classify_chemo("oxaliplatin") == "白金製剤"
    assert mf.classify_chemo("L-OHP, PTX") == "複数/混合"
    assert mf.classify_chemo("パクリタキセル") == "タキサン系"
    assert mf.classify_chemo("") == "不明"


def test_add_study_same_position_on_all_sheets():
    wb, _ = make_wb()
    rows = []
    for n in ("4-5_Claude下書き", "4-5_甲", "4-5_乙"):
        assert mf.add_study_all_blocks(wb[n], "333", "Oki et al. 2015", "RCT") == 2   # 2アウトカム
        rows.append([r["row"] for b in mf.read_45(wb[n]) for r in b["rows"] if r["key"] == "333"])
    assert rows[0] == rows[1] == rows[2]
    assert mf.add_study_all_blocks(wb["4-5_甲"], "333", "Oki et al. 2015") == 0       # 二重追加しない


def test_reconcile_agreement_mismatch_and_final():
    wb, n_rows = make_wb()
    ra, rb = wb["4-5_甲"], wb["4-5_乙"]
    r1 = mf.find_row(ra, "O:A", "111")
    col = mf.C_ITEM0                      # ランダム化
    ra.cell(row=r1, column=col, value=0); rb.cell(row=mf.find_row(rb, "O:A", "111"), column=col, value=0)           # 一致
    ra.cell(row=r1, column=col + 1, value=-1); rb.cell(row=mf.find_row(rb, "O:A", "111"), column=col + 1, value=-2)  # 不一致
    ra.cell(row=r1, column=col + 2, value=-2); rb.cell(row=mf.find_row(rb, "O:A", "111"), column=col + 2, value=-1)  # 不一致→確定で解決
    # 確定列(位置: ブロック0, 行0, 項目2)
    wb["4-5_照合"].cell(row=2 + (0 * n_rows + 0) * mf.N_ITEMS + 2, column=7, value=-2)
    rows, unresolved = mg.reconcile_blocks(wb, "4-5_甲", "4-5_乙")
    s111 = next(r for r in rows if r["key"] == "111" and r["outcome_id"] == "O:A")
    assert s111["items"]["randomization"] == 0
    assert s111["items"]["concealment"] is None and "コンシールメント" in s111["unresolved"]
    assert s111["items"]["blinding_participants"] == -2
    assert unresolved == 1


def test_draft_is_provisional_until_reviewers_enter():
    wb, _ = make_wb()
    d = wb["4-5_Claude下書き"]
    r = mf.find_row(d, "O:A", "222")
    d.cell(row=r, column=mf.C_ITEM0, value=-1)
    d.cell(row=r, column=mf.C_EFF_TYPE, value="RR"); d.cell(row=r, column=mf.C_EFF_VAL, value=0.5); d.cell(row=r, column=mf.C_EFF_CI, value="0.3 to 0.8")
    rows, _ = mg.reconcile_blocks(wb, "4-5_甲", "4-5_乙")
    s = next(x for x in rows if x["key"] == "222" and x["outcome_id"] == "O:A")
    assert s["provisional"] and s["has_data"] and s["items"]["randomization"] == -1
    assert s["effect"] == {"measure": "RR", "point": 0.5, "ci_low": 0.3, "ci_high": 0.8}
    # 評価者が入力したら下書きは使わない
    ra = wb["4-5_甲"]
    ra.cell(row=mf.find_row(ra, "O:A", "222"), column=mf.C_ITEM0, value=0)
    rows, _ = mg.reconcile_blocks(wb, "4-5_甲", "4-5_乙")
    s = next(x for x in rows if x["key"] == "222" and x["outcome_id"] == "O:A")
    assert not s["provisional"]


def test_effect_from_counts_and_ci_parse():
    row = {"effect": {"type": None, "value": None, "ci": None}, "counts": {"ctrl_den": 45, "ctrl_num": 34, "int_den": 44, "int_num": 15}}
    e = mg.build_effect(row, None)
    assert e["measure"] == "RR" and abs(e["point"] - 0.45) < 0.02
    assert mg.parse_ci("-0.68 to 4.41") == (-0.68, 4.41) and mg.parse_ci("0.4–0.9") == (0.4, 0.9)


def test_sr8_read_and_strata():
    wb, _ = make_wb()
    ws = wb["SR-8_エビデンス総体"]
    ws.cell(row=9, column=mf.SR8_COL["bias"], value=-1)
    ws.cell(row=9, column=mf.SR8_COL["strength"], value="B")
    bodies, strata = mf.read_sr8(ws)
    assert [b["outcome_id"] for b in bodies] == ["O:A", "O:B"] and len(strata) == 2
    assert bodies[0]["certainty"] == "B" and bodies[0]["bias"] == -1 and bodies[0]["importance"] == 9
