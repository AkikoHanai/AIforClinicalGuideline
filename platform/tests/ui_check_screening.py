"""スクリーニング管理シートのブラウザ操作テスト(playwright)。
  WS=<作業ディレクトリ> OUT=<出力フォルダ> PLATFORM=<platformのパス> python3 platform/tests/ui_check_screening.py
検証: 検索結果の取り込み、一次・二次の判定、判定の変更(理由が必須)、文献の追加、全文の入手・読了、作業の絞り込み、JSON書き出し、
2名の照合(不一致の検出、協議の結論の反映)、履歴と進捗の出力。"""
import json, os, subprocess, sys
from playwright.sync_api import sync_playwright
import openpyxl
WS = os.environ["WS"]; OUT = os.environ["OUT"]; PL = os.environ["PLATFORM"]
CH = os.environ.get("CHROME", "/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
CQ = "CQ1-牛車腎気丸"; res = []
def ok(n, c, i=""):
    res.append(bool(c)); print(("PASS" if c else "FAIL"), n, i)
subprocess.run([sys.executable, f"{PL}/core/render_screening.py", WS, "-o", f"{WS}/_screening"], check=True, capture_output=True)
url = "file://" + f"{WS}/_screening/{CQ}.screening.html"
def session(b, name, plan):
    ctx = b.new_context(accept_downloads=True); pg = ctx.new_page(); texts = []
    pg.on("dialog", lambda d: (texts.append(d.message), d.accept("理由: " + plan.get("why", "テスト")) if d.type == "prompt" else d.accept()))
    pg.goto(url); pg.fill("#who", name); n0 = pg.evaluate("S.papers.length")
    pg.evaluate("document.getElementById('i_text').value='99000001\\tPaper X 2026\\n99000002\\tPaper Y 2026\\n99000003\\tPaper Z 2026\\n99000001\\tdup'; importText()")
    return ctx, pg, texts, n0
with sync_playwright() as p:
    b = p.chromium.launch(executable_path=CH)
    # ---- 担当A
    ctx, pg, texts, n0 = session(b, "担当A", {"why": "全文で対象外と判断"})
    ok("S1 検索結果の取り込み(重複は追加しない)", pg.evaluate("S.papers.length") == n0 + 3, f"{n0}→{pg.evaluate('S.papers.length')}")
    pg.evaluate("setDec('99000001','s1','採用'); setDec('99000002','s1','除外'); setDec('99000003','s1','保留')")
    ok("S2 一次の判定", pg.evaluate("P('99000001').s1.decision") == "採用" and pg.evaluate("P('99000003').s1.decision") == "保留")
    ok("S3 一次の除外は理由未記入として一覧に出る", pg.evaluate("noReason(P('99000002'))") is True)
    pg.evaluate("setReason('99000002','s1','CIPNに関する研究でない')")
    ok("S4 理由を記入すると理由未記入でなくなる", pg.evaluate("noReason(P('99000002'))") is False)
    pg.evaluate("setDec('99000002','s1','採用')")
    ok("S5 判定の変更は理由が必要(入力を求める)", any("変更" in t for t in texts))
    ok("S6 変更の履歴に理由が残る", any(h["key"] == "99000002" and h["reason"].startswith("理由: ") for h in pg.evaluate("S.history")))
    pg.evaluate("setDec('99000002','s1','除外'); setReason('99000002','s1','CIPNに関する研究でない')")
    pg.evaluate("setDec('99000001','s2','採用'); setFt('99000001','読了'); setFt('99000003','入手済')")
    ok("S7 二次の判定・全文の読了", pg.evaluate("P('99000001').s2.decision") == "採用" and pg.evaluate("P('99000001').ft") == "読了")
    ok("S8 二次の対象で全文が入手済(未読)の絞り込み", pg.evaluate("S.papers.filter(FILTERS.ft1[1]).length") == 1)
    pg.evaluate("document.getElementById('a_pmid').value='99000009'; document.getElementById('a_title').value='Hand search paper'; document.getElementById('a_direct').checked=true; document.getElementById('a_why').value='ASCO 2020の引用文献'; addPaper()")
    ok("S9 文献を二次の対象として追加(一次を経ない)", pg.evaluate("P('99000009').direct") is True and pg.evaluate("need2(P('99000009'))") is True)
    ok("S10 追加した文献が絞り込みに出る", pg.evaluate("S.papers.filter(FILTERS.added[1]).length") == 1)
    with pg.expect_download() as dl: pg.click("text=自分の判定をJSONで書き出す")
    os.makedirs(f"{OUT}/{CQ}", exist_ok=True); dl.value.save_as(f"{OUT}/{CQ}/{CQ}_担当A.screening.json")
    pg.reload(); ok("S11 再読込後も入力が残る", pg.evaluate("P('99000001').s2.decision") == "採用")
    ctx.close()
    # ---- 担当B(99000001の一次を不一致にする)
    ctx, pg, texts, n0 = session(b, "担当B", {"why": "抄録でRCTでない"})
    pg.evaluate("setDec('99000001','s1','除外'); setReason('99000001','s1','研究デザインが対象外(総説・症例報告など)'); setDec('99000002','s1','除外'); setReason('99000002','s1','CIPNに関する研究でない'); setDec('99000003','s1','採用')")
    pg.evaluate("document.getElementById('a_pmid').value='99000009'; document.getElementById('a_title').value='Hand search paper'; document.getElementById('a_direct').checked=true; document.getElementById('a_why').value='ASCO 2020の引用文献'; addPaper(); setDec('99000009','s2','採用')")
    with pg.expect_download() as dl: pg.click("text=自分の判定をJSONで書き出す")
    dl.value.save_as(f"{OUT}/{CQ}/{CQ}_担当B.screening.json"); ctx.close(); b.close()
r = subprocess.run([sys.executable, f"{PL}/core/merge_screening.py", WS, OUT], capture_output=True, text=True); print(r.stdout.strip())
wb = openpyxl.load_workbook(f"{WS}/{CQ}/minds_review.xlsx")
ok("M1 照合・履歴・進捗のシートができる", all(s in wb.sheetnames for s in ["スクリーニング照合", "スクリーニング履歴", "スクリーニング進捗", "スクリーニングログ"]))
rows = {str(r[0]): r for r in wb["スクリーニング照合"].iter_rows(min_row=2, values_only=True) if r[0]}
ok("M2 不一致を検出(99000001の一次)", rows["99000001"][5] == "不一致", str(rows["99000001"][3:6]))
ok("M3 一致を確認(99000002の一次は除外で一致)", rows["99000002"][5] == "一致")
log = {str(r[0]): r for r in wb["スクリーニングログ"].iter_rows(min_row=2, values_only=True) if r[0]}
ok("M4 不一致は最終の判定で「未確定(協議)」", log["99000001"][3] == "未確定(協議)")
ok("M5 一致した除外の理由がログに入る", log["99000002"][3] == "除外" and "CIPN" in (log["99000002"][4] or ""))
ok("M6 追加した文献が反映される", "99000009" in log and "追加: 担当A" in (log["99000009"][7] or ""))
h = list(wb["スクリーニング履歴"].iter_rows(min_row=2, values_only=True))
ok("M7 履歴に変更の理由が残る", any(x[2] == "99000002" and str(x[6]).startswith("理由: ") for x in h), f"{len(h)}件")
# 協議の結論を記入して再照合
ws = wb["スクリーニング照合"]; hdr = [c.value for c in ws[1]]; ci = hdr.index("一次の協議の結論(委員記入)") + 1
for row in ws.iter_rows(min_row=2):
    if str(row[0].value) == "99000001": row[ci - 1].value = "採用"
wb.save(f"{WS}/{CQ}/minds_review.xlsx")
subprocess.run([sys.executable, f"{PL}/core/merge_screening.py", WS, OUT], capture_output=True)
wb = openpyxl.load_workbook(f"{WS}/{CQ}/minds_review.xlsx"); log = {str(r[0]): r for r in wb["スクリーニングログ"].iter_rows(min_row=2, values_only=True) if r[0]}
ok("M8 協議の結論が最終の判定に反映される", log["99000001"][3] == "採用")
prog = list(wb["スクリーニング進捗"].iter_rows(min_row=2, values_only=True))
ok("M9 担当者別の進捗が出る", len(prog) == 2, str([x[0] for x in prog]))
print(sum(res), "/", len(res)); sys.exit(0 if all(res) else 1)
