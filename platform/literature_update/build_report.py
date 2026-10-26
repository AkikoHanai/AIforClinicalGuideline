"""records.json(Methodを読んだ判定案)から、CQ別の集計表とPMID一覧を作る。判定は下書きであり、確定は委員が行う。
  python3 build_report.py  →  out/新規論文_題名限定RCT_Method判定.xlsx / .md"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
R = json.load(open("records.json", encoding="utf-8"))
SEARCH_DATE = "2026-10-26"
def dec(r): return "適格" if r["decision"].startswith("適格") else "要確認" if r["decision"].startswith("要確認") else "不適格"
def scope(r):
    cq, p = r["cq"], r["purpose"]
    if cq.startswith("デュロキセチン") and p == "予防": return "範囲外(企画書のデュロキセチンは治療のCQ)"
    if cq.startswith("ビタミンB12") and p == "予防": return "範囲外(企画書のFRQは治療)"
    if cq.startswith("冷却・圧迫") and p == "治療": return "範囲外(企画書の冷却・圧迫は予防のCQ)"
    if cq.startswith("冷却・圧迫、運動") : return "確認(圧迫は予防のCQ。運動の上乗せは群間比較)"
    if cq.startswith("プレガバリン") and r["decision"].startswith("不適格"): return "—"
    return "範囲内"
GROUPS = [("牛車腎気丸(予防)", ["牛車腎気丸"]), ("ガバペンチノイド", ["プレガバリン(ガバペンチノイド)"]), ("冷却・圧迫(予防)", ["冷却・圧迫", "冷却・圧迫、運動"]),
          ("運動(予防)", ["運動(予防)", "運動"]), ("運動(治療)", ["運動(治療)"]), ("鍼灸(予防)", ["鍼灸(予防)"]), ("鍼灸(治療)", ["鍼灸(治療)"]),
          ("デュロキセチン", ["デュロキセチン"]), ("FRQ ビタミンB12", ["ビタミンB12(FRQ)"])]
EMPTY = ["ミロガバリン", "カルニチン", "アミトリプチリン", "FRQ NSAIDs", "FRQ オピオイド", "FRQ 併用療法"]
rows = []; md = ["# 新規論文の確認(題名に介入語を含むRCT)\n", f"検索日 {SEARCH_DATE}。出典: PubMed。判定は、各論文の抄録のMethodを読んだ下書きで、確定は委員が行う。\n",
    "検索: (\"chemotherapy-induced peripheral neuropathy\"[tiab] OR CIPN[tiab] OR (neuropath*[tiab] AND (chemotherap*[tiab] OR oxaliplatin[tiab] OR paclitaxel[tiab] OR docetaxel[tiab] OR cisplatin[tiab] OR bortezomib[tiab] OR vincristine[tiab] OR taxane*[tiab]))) AND (介入語[ti]) AND (起点[dp] : 2026/10/26[dp]) AND \"randomized controlled trial\"[pt]\n",
    "起点: 牛車腎気丸・カルニチン・アミトリプチリン・ビタミンB12・NSAIDs・オピオイド・鍼灸 2021/07、プレガバリン・ミロガバリン・冷却・圧迫 2021/09、運動 2021/10、デュロキセチン 2022/01。\n",
    "\n## 集計\n", "| CQ | RCT | 適格 | 要確認 | 不適格 | 適格のPMID |\n|---|---|---|---|---|---|"]
summ = []
for g, cqs in GROUPS:
    rs = [r for r in R.values() if r["cq"] in cqs]
    a = [r["pmid"] for r in rs if dec(r) == "適格"]; q = [r["pmid"] for r in rs if dec(r) == "要確認"]; n = [r["pmid"] for r in rs if dec(r) == "不適格"]
    summ.append([g, len(rs), len(a), len(q), len(n), " ".join(a), " ".join(q), " ".join(n)])
    md.append(f"| {g} | {len(rs)} | {len(a)} | {len(q)} | {len(n)} | {' '.join(a)} |")
for g in EMPTY:
    summ.append([g, 0, 0, 0, 0, "", "", ""]); md.append(f"| {g} | 0 | 0 | 0 | 0 | |")
md.append("\n注: 運動(予防)の1件(40750479)は、主要評価が生活の質の副次解析で不適格。冷却・圧迫と運動の併用(41443031、39842054)は、運動に含まれず、冷却・圧迫の行にある。\n")
md.append("## 判定の詳細\n\n| PMID | 年 | 第一著者 | デザイン | 目的 | 化学療法 | N | 介入 | 対照 | CQ | 判定 | 範囲 | Methodの根拠 |\n|---|---|---|---|---|---|---|---|---|---|---|---|---|")
for r in sorted(R.values(), key=lambda x: (x["cq"], -int(x["year"]))):
    md.append("| " + " | ".join([r["pmid"], r["year"], r["author"], r["design"], r["purpose"], r["chemo"], r["n"], r["intervention"], r["comparator"], r["cq"], r["decision"], scope(r), r["basis"]]) + " |")
os.makedirs("out", exist_ok=True); open("out/新規論文_題名限定RCT_Method判定.md", "w", encoding="utf-8").write("\n".join(md) + "\n")
wb = Workbook(); thin = Side(style="thin", color="000000"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
def sheet(ws, head, data, widths):
    ws.append(head)
    for r in data: ws.append(r)
    for row in ws.iter_rows():
        for c in row: c.border = BD; c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9D9D9")
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
sheet(wb.active, ["CQ", "RCT", "適格", "要確認", "不適格", "適格のPMID", "要確認のPMID", "不適格のPMID"], summ, [24, 6, 6, 8, 8, 44, 44, 44]); wb.active.title = "集計"
keys = ["pmid", "year", "author", "design", "purpose", "chemo", "n", "intervention", "comparator", "cq", "decision"]
sheet(wb.create_sheet("判定一覧"), ["PMID", "年", "第一著者", "デザイン", "目的(Method)", "化学療法", "N", "介入", "対照", "CQ", "判定(下書き)", "範囲", "Methodの根拠"],
      [[r[k] for k in keys] + [scope(r), r["basis"]] for r in sorted(R.values(), key=lambda x: (x["cq"], -int(x["year"])))], [10, 6, 14, 22, 8, 22, 7, 26, 22, 18, 36, 28, 60])
wb.save("out/新規論文_題名限定RCT_Method判定.xlsx"); print("ok", len(R))
